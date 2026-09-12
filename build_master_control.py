"""
build_master_control.py — собирает data/master_control.xlsx: единый файл-пульт
для управления ОБОИМИ маркетплейсами (Ozon и WB) сразу. По каждому артикулу
(существующему на любой из площадок, новому черновику на любой из них, или
и там и там) в одной строке видно фото, названия (общее для Ozon и отдельное
для WB — там жёсткий лимит 60 символов), цены, описание, хэштеги и заметки.
Это ЕДИНСТВЕННЫЙ файл, из которого вы редактируете эти поля — после правок
запустите sync-master-control, чтобы перенести их обратно в:
  - data/ozon_catalog.xlsx / data/ozon_new_products.xlsx (Ozon)
  - data/wb_catalog.xlsx / data/wb_new_products.xlsx (WB)
(и, если заменили картинку в столбце "Фото", в новый файл-обложку в
photos/), см. sync_master_control.py.

Сопоставление товаров между площадками — по артикулу (offer_id у Ozon =
vendorCode у WB): если совпадает — это одна строка с обоими статусами;
если артикул WB не совпал ТОЧНО ни с одним Ozon-артикулом (например,
объединённая карточка "0AM325091E / 0AM325091F" или просто по-другому
написанный артикул) — заводится отдельная строка "только WB", чтобы
случайно не перепутать товары; при необходимости сведите такие вручную,
поправив артикул в исходном файле WB.

Название для WB (колонка "Название WB") для товаров, у которых есть карточка
на Ozon, ВСЕГДА строится из названия Ozon (обрезается под лимит WB в 60
символов, с сохранением обозначений коробки/мехатроника DSG/DQ и артикульных
кодов — см. _wb_title) — то, что сейчас реально стоит на WB, при следующем
push-wb-cards будет заменено на это. Для строк "только WB" (артикул не
совпал ни с одним Ozon-товаром) название остаётся как есть на WB.

Цена WB (колонка "Цена WB, ₽") — это ФИНАЛЬНАЯ цена, которую видит
покупатель. Для новых товаров-черновиков это просто цена товара. Для уже
существующих на WB товаров сюда подставляется текущая цена, подтянутая с
WB при последнем fetch-wb; впишете другое число — sync-master-control
перенесёт его в data/wb_catalog.xlsx, а push-wb-price (см. main.py)
отправит на WB, автоматически пересчитав "цену до скидки" так, чтобы
текущий % скидки не изменился (см. wb_catalog_editor.price_before_discount).

master_control.xlsx можно пересобирать сколько угодно раз (build-master-control)
— название/цена/описание/хэштеги/заметки при этом БЕРУТСЯ из исходных файлов
конвейера (то есть уже сделанные там правки не потеряются), а столбец "Фото"
каждый раз строится заново из текущего содержимого photos/.
"""
import logging
import os
import re

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage

import catalog_editor

logger = logging.getLogger("marketplace-agent.build_master_control")

ROOT = os.path.dirname(os.path.abspath(__file__))
PHOTOS_DIR = os.path.join(ROOT, "photos")
CATALOG_PATH = os.path.join(ROOT, "data", "ozon_catalog.xlsx")
NEW_PATH = os.path.join(ROOT, "data", "ozon_new_products.xlsx")
WB_CATALOG_PATH = os.path.join(ROOT, "data", "wb_catalog.xlsx")
WB_NEW_PATH = os.path.join(ROOT, "data", "wb_new_products.xlsx")
MASTER_PATH = os.path.join(ROOT, "data", "master_control.xlsx")

IMAGE_EXTS = {".jpg", ".jpeg", ".png"}
VIDEO_EXTS = {".mp4", ".mov"}
THUMB_DIR = "/tmp/_thumbs_master"
THUMB_PX = 140

STATUS_EXISTING = "Действующий"
STATUS_DRAFT = "Новый (черновик)"

# Ручное сопоставление: артикул WB (vendorCode) -> offer_id на Ozon, для тех
# редких случаев, когда это один и тот же товар, но код на WB отличается от
# Ozon не опечаткой "почти как", а буквально другой строкой (пробел/суффикс
# убран и т.п.) — так что автоматический точный матч (см. run()) их не видит.
# Каждая пара проверена вручную по СОДЕРЖИМОМУ карточек (не только по
# похожести кода!), поэтому список короткий и расширять его нужно так же
# аккуратно — см. sync_master_control.py, там есть обратная версия этого
# словаря (используется, чтобы найти строку в wb_catalog.xlsx по её
# настоящему vendorCode при обратной синхронизации).
WB_ALIAS = {
    # WB "0am325477" — описание слово-в-слово совпадает с Ozon "0am325477ae"
    "0am325477": "0am325477ae",
    # WB "02E305045E" ("Корпус фильтра DQ250 DSG6", просто корпус) — это
    # тот же товар, что Ozon "02e305045" ("Корпус масляного фильтра...",
    # тоже просто корпус, тот же акцент на теплоотводе в тексте). НЕ путать
    # с "02E 305 045" (это на Ozon набор корпус+фильтр вместе) и с
    # "02E305051C" (это отдельно масляный фильтр, не корпус).
    "02E305045E": "02e305045",
}

HEADERS = [
    "Артикул (offer_id)",
    "Статус Ozon",
    "Статус WB",
    "Фото",
    "Кол-во файлов фото/видео",
    "Название товара",
    "Название WB (до 60 симв.)",
    "Цена, ₽",
    "Цена до скидки, ₽",
    "Остаток, шт.",
    "Цена WB, ₽ (финальная, покупателю)",
    "Описание",
    "Хэштеги / Теги",
    "Заметки",
    "Файлы (папка photos/)",
]
COL_WIDTHS = [22, 16, 16, 20, 10, 34, 24, 12, 14, 12, 14, 50, 30, 34, 40]


def _find_header(headers, *, must_contain, must_not_contain=()):
    for i, x in enumerate(headers):
        if not x:
            continue
        if must_contain not in x:
            continue
        if any(bad in x for bad in must_not_contain):
            continue
        return i
    return None


def _read_existing():
    """offer_id -> {name, price, old_price, quantity, description, notes, tnved, hashtags, photo_url}."""
    if not os.path.exists(CATALOG_PATH):
        return {}
    wb = openpyxl.load_workbook(CATALOG_PATH)
    ws = wb.active
    h = [c.value for c in ws[1]]
    idx_name = h.index("Название товара")
    idx_desc = h.index("Описание")
    idx_notes = h.index("Заметки")
    idx_tnved = [i for i, x in enumerate(h) if x and "ТН ВЭД" in x][0]
    idx_photo = [i for i, x in enumerate(h) if x and x.startswith("Фото")][0]
    idx_hashtags = next((i for i, x in enumerate(h) if x and "Хэштег" in x), None)
    idx_qty = next((i for i, x in enumerate(h) if x and "Кол-во к продаже" in x), None)
    idx_price = _find_header(h, must_contain="Цена", must_not_contain=("до скидки",))
    idx_old_price = _find_header(h, must_contain="до скидки")

    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        oid = str(row[0]).strip()
        photo_cell = row[idx_photo] or ""
        first_url = photo_cell.split("|")[0].strip() if photo_cell else ""
        out[oid] = {
            "name": row[idx_name] or "",
            "price": row[idx_price] if idx_price is not None else None,
            "old_price": row[idx_old_price] if idx_old_price is not None else None,
            "quantity": row[idx_qty] if idx_qty is not None else None,
            "description": row[idx_desc] or "",
            "notes": row[idx_notes] or "",
            "tnved": row[idx_tnved] or "",
            "hashtags": (row[idx_hashtags] or "") if idx_hashtags is not None else "",
            "photo_url": first_url,
        }
    return out


def _read_draft():
    """offer_id -> {name, price, old_price, quantity, description, notes, sample, hashtags}."""
    if not os.path.exists(NEW_PATH):
        return {}
    wb = openpyxl.load_workbook(NEW_PATH)
    ws = wb.active
    h = [c.value for c in ws[1]]
    idx_name = next(i for i, x in enumerate(h) if x and "Название" in x)
    idx_desc = next(i for i, x in enumerate(h) if x and "Описание" in x)
    idx_notes = next(i for i, x in enumerate(h) if x and "Заметки" in x)
    idx_sample = next(i for i, x in enumerate(h) if x and "Образец" in x)
    idx_hashtags = next((i for i, x in enumerate(h) if x and "Хэштег" in x), None)
    idx_qty = next((i for i, x in enumerate(h) if x and "Кол-во к продаже" in x), None)
    idx_price = _find_header(h, must_contain="Цена", must_not_contain=("до скидки",))
    idx_old_price = _find_header(h, must_contain="до скидки")

    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        oid = str(row[0]).strip()
        out[oid] = {
            "name": row[idx_name] or "",
            "price": row[idx_price] if idx_price is not None else None,
            "old_price": row[idx_old_price] if idx_old_price is not None else None,
            "quantity": row[idx_qty] if idx_qty is not None else None,
            "description": row[idx_desc] or "",
            "notes": row[idx_notes] or "",
            "sample": row[idx_sample] or "",
            "hashtags": (row[idx_hashtags] or "") if idx_hashtags is not None else "",
        }
    return out


def _read_wb_existing():
    """vendorCode -> {title, description, notes, price, discount}."""
    if not os.path.exists(WB_CATALOG_PATH):
        return {}
    wb = openpyxl.load_workbook(WB_CATALOG_PATH)
    ws = wb.active
    h = [c.value for c in ws[1]]
    idx_title = next(i for i, x in enumerate(h) if x and "Название" in x)
    idx_desc = next(i for i, x in enumerate(h) if x and "Описание" in x)
    idx_notes = next((i for i, x in enumerate(h) if x and "Заметки" in x), None)
    idx_price = next((i for i, x in enumerate(h) if x and x.startswith("Цена WB")), None)
    idx_discount = next((i for i, x in enumerate(h) if x and x.startswith("Скидка WB")), None)

    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        vc = str(row[0]).strip()
        out[vc] = {
            "title": row[idx_title] or "",
            "description": row[idx_desc] or "",
            "notes": (row[idx_notes] or "") if idx_notes is not None else "",
            "price": row[idx_price] if idx_price is not None else None,
            "discount": row[idx_discount] if idx_discount is not None else None,
        }
    return out


def _read_wb_draft():
    """vendorCode -> {title, description, price, notes, sample}."""
    if not os.path.exists(WB_NEW_PATH):
        return {}
    wb = openpyxl.load_workbook(WB_NEW_PATH)
    ws = wb.active
    h = [c.value for c in ws[1]]
    idx_title = next(i for i, x in enumerate(h) if x and "Название" in x)
    idx_desc = next(i for i, x in enumerate(h) if x and "Описание" in x)
    idx_notes = next((i for i, x in enumerate(h) if x and "Заметки" in x), None)
    idx_sample = next((i for i, x in enumerate(h) if x and "Образец" in x), None)
    idx_price = _find_header(h, must_contain="Цена", must_not_contain=("до скидки",))

    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        vc = str(row[0]).strip()
        out[vc] = {
            "title": row[idx_title] or "",
            "description": row[idx_desc] or "",
            "price": row[idx_price] if idx_price is not None else None,
            "notes": (row[idx_notes] or "") if idx_notes is not None else "",
            "sample": (row[idx_sample] or "") if idx_sample is not None else "",
        }
    return out


_CODE_RE = re.compile(r"^(DSG\d|(?:DQ|DL|VL)\d{3})$", re.IGNORECASE)
_PARTNO_RE = re.compile(r"^0[A-Za-z0-9]{2,9}$")


def _is_protected_word(word, offer_id):
    """
    Слово в названии, которое нельзя терять при обрезке под лимит WB (60
    симв.): обозначение коробки/мехатроника (DSG6/DSG7, DQ200/DQ250/DQ500 и
    т.п.), артикульные коды VAG-образца (обычно начинаются с "0": 0AM, 0CW,
    0B5, 02E...) и сам артикул товара — эти вещи для покупателя в этой нише
    важнее любых других слов в названии.
    """
    w = word.strip(".,;:()[]/\"'")
    if not w:
        return False
    if _CODE_RE.match(w):
        return True
    if _PARTNO_RE.match(w):
        return True
    if offer_id and w.lower() == str(offer_id).lower():
        return True
    return False


_STRIP_CHARS = ".,;:()[]/\"'"


def _wb_title(name, offer_id, limit=60):
    """
    Строит короткое название для WB (лимит 60 симв.) из полного названия
    Ozon: если оно и так укладывается — используется как есть; если нет —
    "защищённые" слова (коды коробки/мехатроника/артикул, см.
    _is_protected_word) сохраняются ВСЕГДА. Если в названии есть перечисление
    через "+" (частый случай для комплектов: "... + 2 пыльника + 2 крышки")
    — обрезаем целыми пунктами через "+" (см. _wb_title_by_segments), чтобы
    никогда не оставить "повисшую" цифру или плюс без слова, к которому они
    относятся. Если "+" в названии нет — обрезаем по словам (см.
    _wb_title_by_words).
    """
    name = (name or "").strip()
    if not name:
        return ""
    if len(name) <= limit:
        return name
    if "+" in name:
        return _wb_title_by_segments(name, offer_id, limit)
    return _wb_title_by_words(name, offer_id, limit)


def _wb_title_by_words(name, offer_id, limit):
    """
    Обрезка по отдельным словам: "защищённые" слова (коды коробки/
    мехатроника/артикул) сохраняются ВСЕГДА, остальные добавляются по
    порядку слева направо, пока хватает места. Порядок слов в результате
    всегда совпадает с порядком в исходном названии. Со слов снимается
    обрамляющая пунктуация (скобки, запятые) — иначе может остаться
    "повисшая" закрывающая скобка без открывающей.
    """
    words = [w.strip(_STRIP_CHARS) for w in name.split()]
    words = [w for w in words if w]
    protected_idx = [i for i, w in enumerate(words) if _is_protected_word(w, offer_id)]
    must_words = [words[i] for i in protected_idx]
    must_len = sum(len(w) for w in must_words) + max(0, len(must_words) - 1)

    if must_len > limit:
        # даже "защищённые" слова целиком не помещаются — берём их по
        # порядку, сколько влезет, дальше жёстко обрезаем по длине.
        acc, cur = [], 0
        for w in must_words:
            add = len(w) + (1 if acc else 0)
            if cur + add > limit:
                break
            acc.append(w)
            cur += add
        return " ".join(acc)[:limit]

    budget = limit - must_len
    selected = set(protected_idx)
    used = 0
    for i, w in enumerate(words):
        if i in selected:
            continue
        add = len(w) + 1  # +1 — примерно на разделяющий пробел
        if used + add <= budget:
            selected.add(i)
            used += add

    result = " ".join(words[i] for i in sorted(selected))
    return result[:limit].rstrip() if len(result) > limit else result


def _wb_title_by_segments(name, offer_id, limit):
    """
    Обрезка целыми пунктами перечисления через "+" (например "Ремкомплект
    ... + 2 пыльника + 2 крышки" -> пункты "Ремкомплект ...", "2 пыльника",
    "2 крышки"). Пункт с "защищённым" словом (код коробки/артикул) остаётся
    ВСЕГДА; остальные пункты добавляются целиком, только если помещаются —
    никогда не берётся часть пункта, чтобы не оставить голую цифру или "+"
    без слова, к которому она относится.
    """
    segments = [s.strip() for s in name.split("+")]
    segments = [s for s in segments if s]

    def seg_protected(seg):
        return any(_is_protected_word(w, offer_id) for w in seg.split())

    def join_len(idxs):
        return len(" + ".join(segments[i] for i in idxs))

    protected_idx = [i for i, s in enumerate(segments) if seg_protected(s)]
    if protected_idx and join_len(protected_idx) > limit:
        # даже обязательные пункты целиком не помещаются — обрезаем по словам
        return _wb_title_by_words(name, offer_id, limit)

    selected = set(protected_idx)
    for i in range(len(segments)):
        if i in selected:
            continue
        if join_len(sorted(selected | {i})) <= limit:
            selected.add(i)

    if not selected:
        return _wb_title_by_words(name, offer_id, limit)

    result = " + ".join(segments[i] for i in sorted(selected))
    return result[:limit].rstrip() if len(result) > limit else result


def _make_thumb(src_path, dst_path):
    try:
        img = PILImage.open(src_path)
        img = img.convert("RGB")
        img.thumbnail((THUMB_PX, THUMB_PX))
        img.save(dst_path, "JPEG", quality=80)
        return True
    except Exception as exc:  # noqa: BLE001
        logger.warning("Не удалось построить эскиз %s: %s", src_path, exc)
        return False


def _apply_wb_alias(d):
    """
    Переносит записи из WB_ALIAS под ключ Ozon-артикула, которому они
    реально соответствуют — чтобы такая строка не заводилась отдельно как
    "только WB", а слилась с обычной строкой этого товара. Настоящий
    vendorCode сохраняется в "_wb_vendor_code" — он нужен sync_master_control.py,
    чтобы знать, в какую строку wb_catalog.xlsx/wb_new_products.xlsx реально
    писать (там ключ — оригинальный vendorCode, а не Ozon offer_id).

    Если целевой ключ (Ozon-артикул из WB_ALIAS) САМ по себе уже существует
    как отдельный, настоящий vendorCode на WB — редкое, но реальное
    совпадение строк — слияние для этой пары ОТМЕНЯЕТСЯ: обе записи
    остаются отдельными строками под своими настоящими vendorCode. Иначе
    одна из двух разных карточек молча теряется (перезаписывается другой
    в словаре). В лог пишется предупреждение — такую пару стоит проверить
    вручную и, возможно, убрать из WB_ALIAS (см. docstring WB_ALIAS).
    """
    direct_targets = {vendor_code for vendor_code in d if vendor_code in WB_ALIAS.values()}

    out = {}
    for vendor_code, v in d.items():
        target_oid = WB_ALIAS.get(vendor_code, vendor_code)
        entry = dict(v)
        entry["_wb_vendor_code"] = vendor_code
        if target_oid != vendor_code and target_oid in direct_targets:
            logger.warning(
                "WB_ALIAS: %s -> %s не объединены, потому что на WB уже есть "
                "отдельный настоящий vendorCode %s — похоже, это два разных "
                "товара. Обе строки оставлены отдельно; проверьте вручную.",
                vendor_code, target_oid, target_oid,
            )
            out[vendor_code] = entry
        else:
            out[target_oid] = entry
    return out


def run() -> int:
    os.makedirs(THUMB_DIR, exist_ok=True)
    os.makedirs(os.path.dirname(MASTER_PATH), exist_ok=True)

    existing = _read_existing()
    draft = _read_draft()
    wb_existing = _apply_wb_alias(_read_wb_existing())
    wb_draft = _apply_wb_alias(_read_wb_draft())

    ozon_ids = set(existing) | set(draft)
    # артикул WB считается "тем же товаром", только если ТОЧНО (с учётом
    # регистра) совпадает с уже известным Ozon-артикулом — иначе заводим
    # отдельную строку "только WB", чтобы не перепутать разные товары.
    wb_only_existing = {k: v for k, v in wb_existing.items() if k not in ozon_ids}
    wb_only_draft = {k: v for k, v in wb_draft.items() if k not in ozon_ids and k not in wb_only_existing}
    known_ids = sorted(ozon_ids | set(wb_only_existing) | set(wb_only_draft))

    # Медиафайлы: с 2026-09-10 у каждого товара своя папка photos/<артикул>/
    # — всё внутри неё однозначно принадлежит этому товару, без угадывания
    # по префиксу имени файла (см. catalog_editor.own_media_files). Файлы,
    # ещё не перенесённые в свою папку (старый плоский способ) и файлы с
    # именем, не совпадающим ни с одним известным артикулом, остаются на
    # верхнем уровне photos/ — они и попадают в "Нераспознанные фото".
    groups = {}
    for oid in known_ids:
        files = catalog_editor.own_media_files(PHOTOS_DIR, oid, IMAGE_EXTS | VIDEO_EXTS)
        if files:
            groups[oid] = files

    unmatched = []
    media_files_count = sum(len(v) for v in groups.values())
    if os.path.isdir(PHOTOS_DIR):
        top_level = sorted(
            f for f in os.listdir(PHOTOS_DIR)
            if os.path.isfile(os.path.join(PHOTOS_DIR, f))
            and os.path.splitext(f)[1].lower() in (IMAGE_EXTS | VIDEO_EXTS)
        )
        # Файл на верхнем уровне уже "занят", если own_media_files нашла его
        # там же через старый плоский резервный разбор (это происходит,
        # только пока у товара ещё нет своей папки).
        claimed_flat = {
            f
            for oid, files in groups.items()
            if not os.path.isdir(os.path.join(PHOTOS_DIR, oid))
            for f in files
        }
        unmatched = [f for f in top_level if f not in claimed_flat]
        media_files_count += len(unmatched)

    out = openpyxl.Workbook()
    ws = out.active
    ws.title = "Товары"
    ws.append(HEADERS)
    header_fill = PatternFill(start_color="FFD9D9D9", end_color="FFD9D9D9", fill_type="solid")
    for c in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for i, w in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    thin = Side(style="thin", color="FFBBBBBB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    row_i = 2
    for oid in known_ids:
        is_existing = oid in existing
        is_draft = oid in draft
        status_ozon = STATUS_EXISTING if is_existing else (STATUS_DRAFT if is_draft else "")

        is_wb_existing = oid in wb_existing
        is_wb_draft = oid in wb_draft
        status_wb = STATUS_EXISTING if is_wb_existing else (STATUS_DRAFT if is_wb_draft else "")

        src = existing.get(oid) or draft.get(oid) or {}
        wb_src = wb_existing.get(oid) or wb_draft.get(oid) or {}

        name = src.get("name", "")
        price = src.get("price")
        old_price = src.get("old_price")
        quantity = src.get("quantity")
        # описание общее на обе площадки: берём Ozon-версию, если товара
        # нет на Ozon (только WB) — берём WB-версию.
        desc = src.get("description", "") or wb_src.get("description", "")
        notes = src.get("notes", "") or wb_src.get("notes", "")
        hashtags = src.get("hashtags", "")
        # название для WB: если товар есть на Ozon — ВСЕГДА берём оттуда
        # (обрезая под лимит WB, см. _wb_title), даже если на WB сейчас
        # стоит другое название — оно будет перезаписано при следующем
        # push-wb-cards. Если товара на Ozon нет (строка "только WB") —
        # берём то, что уже стоит на WB, как раньше.
        name_wb = _wb_title(name, oid) if src else wb_src.get("title", "")
        # Цена WB: для черновиков — то, что уже вписано в wb_new_products.xlsx;
        # для уже существующих на WB товаров — текущая финальная цена,
        # подтянутая с WB при последнем fetch-wb (см. wb_catalog_editor.
        # current_wb_price). Впишете сюда другое число — при следующем
        # push-wb-price уйдёт новая цена, текущий % скидки не изменится.
        price_wb = wb_draft.get(oid, {}).get("price")
        if price_wb in (None, "") and is_wb_existing:
            price_wb = wb_existing.get(oid, {}).get("price")

        fgroup = groups.get(oid, [])
        img_files = [f for f in fgroup if os.path.splitext(f)[1].lower() in IMAGE_EXTS]
        file_count = len(fgroup)
        if fgroup:
            files_str = ", ".join(fgroup)
        else:
            url = (existing.get(oid) or {}).get("photo_url", "")
            files_str = (
                "фото уже на карточке Ozon, локального файла нет: " + url
                if url else "(нет фото ни локально, ни на карточке)"
            )

        ws.cell(row=row_i, column=1, value=oid)
        ws.cell(row=row_i, column=2, value=status_ozon)
        ws.cell(row=row_i, column=3, value=status_wb)
        ws.cell(row=row_i, column=5, value=file_count)
        ws.cell(row=row_i, column=6, value=name)
        ws.cell(row=row_i, column=7, value=name_wb)
        ws.cell(row=row_i, column=8, value=price)
        ws.cell(row=row_i, column=9, value=old_price)
        ws.cell(row=row_i, column=10, value=quantity)
        ws.cell(row=row_i, column=11, value=price_wb)
        ws.cell(row=row_i, column=12, value=desc)
        ws.cell(row=row_i, column=13, value=hashtags)
        ws.cell(row=row_i, column=14, value=notes)
        ws.cell(row=row_i, column=15, value=files_str)

        for c in range(1, len(HEADERS) + 1):
            ws.cell(row=row_i, column=c).alignment = Alignment(vertical="top", wrap_text=True)
            ws.cell(row=row_i, column=c).border = border

        if img_files:
            src_path = catalog_editor.media_path(PHOTOS_DIR, oid, img_files[0])
            thumb_path = os.path.join(THUMB_DIR, oid.replace("/", "_") + ".jpg")
            if _make_thumb(src_path, thumb_path):
                xlimg = XLImage(thumb_path)
                xlimg.width = 100
                xlimg.height = 100
                ws.add_image(xlimg, "D" + str(row_i))

        ws.row_dimensions[row_i].height = 78
        row_i += 1

    ws.freeze_panes = "A2"

    ws2 = out.create_sheet("Нераспознанные фото")
    ws2.append(["Файл", "Похожий артикул", "Комментарий"])
    for c in range(1, 4):
        ws2.cell(row=1, column=c).font = Font(bold=True)
        ws2.cell(row=1, column=c).fill = header_fill
    ws2.column_dimensions["A"].width = 30
    ws2.column_dimensions["B"].width = 30
    ws2.column_dimensions["C"].width = 50
    r = 2
    for f in sorted(unmatched):
        base, _ = os.path.splitext(f)
        prefix = base.split("_")[0] if "_" in base else base.split("-")[0]
        ws2.cell(row=r, column=1, value=f)
        ws2.cell(row=r, column=2, value=prefix)
        r += 1

    out.save(MASTER_PATH)
    print(f"Готово: {MASTER_PATH}")
    print(
        f"Известных артикулов: {len(known_ids)} "
        f"(Ozon существующих {len(existing)}, Ozon черновиков {len(draft)}, "
        f"только-WB существующих {len(wb_only_existing)}, только-WB черновиков {len(wb_only_draft)})"
    )
    print(f"Медиафайлов: {media_files_count}, сопоставлено {sum(len(v) for v in groups.values())} в {len(groups)} товарах")
    if unmatched:
        print(f"Нераспознанных файлов: {len(unmatched)} — см. лист 'Нераспознанные фото' в {MASTER_PATH}")
    return 0


if __name__ == "__main__":
    import sys
    sys.exit(run())
