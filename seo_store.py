"""
Хранилище собранной SEO-семантики и готовых текстов карточек —
data/seo_keywords.xlsx, два листа:
  - "Семантика"        — артикул, ключевое слово, частотность, источник, дата
  - "Тексты Ozon-WB"    — готовые названия/описания под публикацию

Файл накопительный: повторный сбор по тому же артикулу не дублирует уже
имеющиеся пары (артикул, фраза) — только добавляет новое.
"""
import datetime
import os
from typing import Dict

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DEFAULT_PATH = os.path.join(os.path.dirname(__file__), "data", "seo_keywords.xlsx")
QUEUE_PATH = os.path.join(os.path.dirname(__file__), "data", "wordstat_queue.xlsx")

SEM_SHEET = "Семантика"
TEXT_SHEET = "Тексты Ozon-WB"
SEM_HEADERS = ["Артикул", "Ключевое слово", "Частотность", "Источник", "Дата добавления"]
TEXT_HEADERS = ["Артикул", "Название (WB, до 60 симв.)", "Название (Ozon)", "Описание", "Дата"]
QUEUE_HEADERS = ["Артикул", "Стартовая фраза (необязательно; через ';' если несколько)"]

_HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="404040")
_BODY_FONT = Font(name="Arial", size=10)


def _write_header(ws, headers) -> None:
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = _HEADER_FONT
        c.fill = _HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _ensure_workbook(path: str) -> "openpyxl.Workbook":
    if os.path.exists(path):
        return openpyxl.load_workbook(path)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SEM_SHEET
    _write_header(ws, SEM_HEADERS)
    ws2 = wb.create_sheet(TEXT_SHEET)
    _write_header(ws2, TEXT_HEADERS)
    return wb


def add_semantics(
    article: str, phrases: Dict[str, int], source: str = "wordstat_api", path: str = DEFAULT_PATH
) -> int:
    """
    Добавляет собранные фразы в лист "Семантика" для артикула, не дублируя
    уже имеющиеся пары (артикул, фраза). Возвращает количество новых строк.
    """
    wb = _ensure_workbook(path)
    if SEM_SHEET not in wb.sheetnames:
        ws = wb.create_sheet(SEM_SHEET, 0)
        _write_header(ws, SEM_HEADERS)
    else:
        ws = wb[SEM_SHEET]

    existing = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0] or not row[1]:
            continue
        existing.add((str(row[0]).strip().casefold(), str(row[1]).strip().casefold()))

    today = datetime.date.today().isoformat()
    added = 0
    r = ws.max_row + 1
    for phrase, count in sorted(phrases.items(), key=lambda kv: -kv[1]):
        key = (article.strip().casefold(), phrase.strip().casefold())
        if key in existing:
            continue
        # ВАЖНО: помечаем ключ добавленным СРАЗУ, а не только сверяемся с тем,
        # что было в файле ДО этого вызова. Без этого две разные по написанию
        # фразы с одинаковым casefold (например "0am325091E" из totalCount и
        # "0am325091e" из results/associations — см. чат) обе проходили бы
        # проверку "key in existing" как "новые" и обе записывались бы в
        # одном и том же вызове — реальный источник дублей, найденный на
        # живых данных.
        existing.add(key)
        ws.cell(row=r, column=1, value=article).font = _BODY_FONT
        ws.cell(row=r, column=2, value=phrase).font = _BODY_FONT
        ws.cell(row=r, column=3, value=count).font = _BODY_FONT
        ws.cell(row=r, column=4, value=source).font = _BODY_FONT
        ws.cell(row=r, column=5, value=today).font = _BODY_FONT
        r += 1
        added += 1

    widths = [16, 55, 22, 16, 16]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(SEM_HEADERS))}{max(r - 1, 1)}"

    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)
    return added


def ensure_queue_template(path: str = QUEUE_PATH) -> str:
    """
    Создаёт (если ещё нет) data/wordstat_queue.xlsx — простой список: колонка
    "Артикул" (по одному в строке, просто вставьте столбец из вашей таблицы)
    и необязательная колонка "Стартовая фраза" (если оставить пустой —
    в качестве стартовой фразы для Wordstat используется сам артикул).
    Возвращает путь к файлу.
    """
    if os.path.exists(path):
        return path
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Очередь"
    _write_header(ws, QUEUE_HEADERS)
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 60
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)
    return path


def read_queue(path: str = QUEUE_PATH) -> list:
    """
    Читает data/wordstat_queue.xlsx и возвращает список (артикул, стартовая_фраза)
    — если стартовая фраза не задана, вместо неё используется сам артикул.
    Строки с пустым артикулом пропускаются.
    """
    if not os.path.exists(path):
        return []
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        article = str(row[0]).strip()
        if not article:
            continue
        seed = str(row[1]).strip() if len(row) > 1 and row[1] else article
        out.append((article, seed))
    return out


def get_semantics_for_article(article: str, path: str = DEFAULT_PATH) -> Dict[str, int]:
    """Уже собранная (в прошлые разы) семантика для артикула, если файл существует."""
    if not os.path.exists(path):
        return {}
    wb = openpyxl.load_workbook(path, data_only=True)
    if SEM_SHEET not in wb.sheetnames:
        return {}
    ws = wb[SEM_SHEET]
    out: Dict[str, int] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0] or not row[1]:
            continue
        if str(row[0]).strip().casefold() != article.strip().casefold():
            continue
        try:
            count = int(row[2] or 0)
        except (TypeError, ValueError):
            count = 0
        phrase = str(row[1]).strip()
        if phrase not in out or count > out[phrase]:
            out[phrase] = count
    return out
