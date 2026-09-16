"""
Учёт комплектов (наборов), собираемых из тех же деталей, что продаются и
по отдельности — например "ремкомплект штоков" = 2х пыльник + 2х сальник +
2х крышка штока, где каждая из этих трёх деталей также продаётся сама по
себе своим артикулом.

У комплекта НЕТ собственного физического остатка — общий склад это
остаток деталей. "Кол-во к продаже" в строке-комплекте в ozon_catalog.xlsx
полностью ВЫЧИСЛЯЕТСЯ: сколько комплектов реально можно собрать из того,
что сейчас есть по каждой входящей детали (минимум по всем деталям).

Состав комплектов задаётся вручную, один раз на комплект, в data/kits.xlsx
(см. main.py --build-kits-template) — по одной строке на каждую деталь,
входящую в комплект.

Встраивается в общий склад (stock_sync.py) в двух местах:
  1. expand_deltas() — перед записью изменений в ozon_catalog.xlsx: если
     среди проданных артикулов оказался сам комплект, его дельта
     заменяется дельтами деталей (продажа 1 комплекта = списание N штук
     каждой детали). У самого комплекта дельта не применяется напрямую —
     его остаток не независимый.
  2. recompute_kit_stock() — после записи: пересчитывает "Кол-во к
     продаже" у строк-комплектов от актуального (уже списанного) остатка
     деталей и сразу пишет обратно в тот же файл — следующий push-stock/
     push-wb-stock отправит это число на маркетплейсы как обычно, никакой
     специальной обработки для комплектов на стороне Ozon/WB не нужно.

Если data/kits.xlsx нет или в нём пусто — все функции ведут себя как
no-op (комплектов просто не заведено, обычный учёт по отдельным деталям
работает как раньше).
"""
import os
from typing import Dict, List, Tuple

import openpyxl

DEFAULT_KITS_PATH = os.path.join(os.path.dirname(__file__), "data", "kits.xlsx")

KITS_HEADER = [
    "Артикул комплекта",
    "Название комплекта (для справки)",
    "Артикул детали",
    "Кол-во детали в комплекте",
]


def load_bom(kits_path: str = DEFAULT_KITS_PATH) -> Dict[str, List[Tuple[str, int]]]:
    """
    kit_offer_id -> [(component_offer_id, qty_per_kit), ...].
    Пустой словарь, если файла нет — значит, комплекты ещё не заведены.
    """
    if not os.path.exists(kits_path):
        return {}
    wb = openpyxl.load_workbook(kits_path)
    ws = wb.active
    bom: Dict[str, List[Tuple[str, int]]] = {}
    for row in ws.iter_rows(min_row=2):
        kit_id = row[0].value
        component_id = row[2].value if len(row) > 2 else None
        qty = row[3].value if len(row) > 3 else None
        if not kit_id or not component_id:
            continue
        kit_id = str(kit_id).strip()
        component_id = str(component_id).strip()
        try:
            qty = int(qty) if qty not in (None, "") else 1
        except (TypeError, ValueError):
            qty = 1
        if qty <= 0:
            qty = 1
        bom.setdefault(kit_id, []).append((component_id, qty))
    return bom


def expand_deltas(deltas: Dict[str, int], bom: Dict[str, List[Tuple[str, int]]]) -> Dict[str, int]:
    """
    Заменяет дельту комплекта дельтами его деталей (см. docstring модуля).
    Артикулы, которые не являются комплектом, переносятся как есть.
    """
    if not bom:
        return dict(deltas)
    expanded: Dict[str, int] = {}
    for offer_id, delta in deltas.items():
        components = bom.get(offer_id)
        if not components:
            expanded[offer_id] = expanded.get(offer_id, 0) + delta
            continue
        # delta отрицательная при продаже (списание), положительная при
        # возврате/отмене — kits_delta = сколько комплектов продано (может
        # быть отрицательным при возврате), переносим на каждую деталь.
        kits_delta = -delta
        for component_id, qty_per_kit in components:
            expanded[component_id] = expanded.get(component_id, 0) - kits_delta * qty_per_kit
    return expanded


def recompute_kit_stock(catalog_path: str, bom: Dict[str, List[Tuple[str, int]]]) -> Dict[str, int]:
    """
    Пересчитывает "Кол-во к продаже" для строк-комплектов = сколько
    комплектов можно собрать из текущего остатка деталей (минимум по всем
    входящим деталям, целое деление). Пишет прямо в catalog_path через
    catalog_editor.apply_stock_deltas (та же функция, что и для обычных
    продаж — остаток не уходит ниже 0, изменённые ячейки подсвечиваются).
    Возвращает {kit_offer_id: новый_остаток} только для реально
    изменившихся комплектов.
    """
    if not bom:
        return {}
    import catalog_editor

    stock = catalog_editor.get_stock_levels(catalog_path)
    kit_deltas: Dict[str, int] = {}
    for kit_id, components in bom.items():
        if kit_id not in stock:
            continue  # комплекта нет строкой в каталоге — нечего обновлять
        possible = [stock.get(component_id, 0) // qty_per_kit for component_id, qty_per_kit in components]
        max_kits = min(possible) if possible else 0
        current = stock.get(kit_id, 0)
        if max_kits != current:
            kit_deltas[kit_id] = max_kits - current

    if not kit_deltas:
        return {}
    new_values, _unmatched = catalog_editor.apply_stock_deltas(catalog_path, kit_deltas)
    return new_values


def build_kits_template(kits_path: str = DEFAULT_KITS_PATH) -> bool:
    """
    Создаёт пустой data/kits.xlsx с заголовками, если его ещё нет.
    True — создал, False — файл уже существовал (ничего не тронуто).
    """
    if os.path.exists(kits_path):
        return False
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Комплекты"
    ws.append(KITS_HEADER)
    for col_idx, width in enumerate([20, 35, 20, 18], start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width
    os.makedirs(os.path.dirname(kits_path), exist_ok=True)
    wb.save(kits_path)
    return True
