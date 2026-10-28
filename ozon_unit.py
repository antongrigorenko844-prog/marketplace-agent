"""
Юнит-экономика Ozon по товарам: недельный алерт в Telegram + причины возвратов. Только чтение — на Ozon ничего не меняет.

Команды (main.py):
  --ozon-returns       возвраты Ozon за 90 дней (/v1/returns/list): товар, причина, статус →
                       data/ozon_returns.json (+ сводка «товар → причины»).
  --ozon-unit-weekly   раз в неделю (понедельник, workflow ozon-unit.yml):
                       1) свежие начисления Ozon за 30 дней (ozon-finance, с разбивкой услуг по товару);
                       2) возвраты за 90 дней;
                       3) прибыль товара за МЕСЯЦ = начисления Ozon − налог 6% с выручки − закупка × (продано − возвращено);
                       4) в Telegram владельцу: товары в минусе за месяц (не за неделю — расходы Ozon списывает сразу,
                          а выручку при выкупе, недельный минус почти всегда ложный), что съело деньги, причины возвратов;
                          отдельно — дешёвые товары, где «Звёздные товары»/сбор отзывов съедают >20% выручки.
                       Таблица целиком — data/ozon_unit_report.xlsx.
"""
from __future__ import annotations

import json
import os
from collections import Counter, defaultdict
from datetime import datetime, timedelta, timezone

DATA = "data"
RETURNS_PATH = os.path.join(DATA, "ozon_returns.json")
REPORT_PATH = os.path.join(DATA, "ozon_unit_report.xlsx")
STATE_PATH = os.path.join(DATA, "ozon_unit_state.json")
PAID_PROGRAMS = ("Звёздные товары", "Сбор первых отзывов", "Ускоренный сбор отзывов", "Продвижение бренда")
PAID_SHARE_ALERT = 0.20      # платные программы съедают больше 20% выручки товара — сказать владельцу
CHEAP_PRICE = 5000           # «дешёвый» товар — цена на Ozon ниже этой


# ---------------------------------------------------------------- возвраты
def fetch_returns(days: int = 90) -> dict:
    import ozon_client
    to = datetime.now(timezone.utc)
    frm = to - timedelta(days=days)
    iso = lambda d: d.strftime("%Y-%m-%dT%H:%M:%S.000Z")  # noqa: E731
    filters = [{"logistic_return_date": {"time_from": iso(frm), "time_to": iso(to)}},
               {"visual_status_change_moment": {"time_from": iso(frm), "time_to": iso(to)}},
               {}]
    rows, errors, used = [], [], None
    for flt in filters:
        rows, last_id = [], 0
        try:
            for _ in range(40):
                res = ozon_client._post("/v1/returns/list", {"filter": flt, "limit": 500, "last_id": last_id})
                part = res.get("returns") or []
                rows += part
                if not res.get("has_next") or not part:
                    break
                last_id = part[-1].get("id") or 0
            used = flt
            break
        except Exception as exc:  # noqa: BLE001
            errors.append(f"фильтр {list(flt) or 'без фильтра'}: {exc}")
    out = []
    for r in rows:
        pr = r.get("product") or {}
        st = ((r.get("visual") or {}).get("status") or {}).get("display_name") or r.get("status") or ""
        lg = r.get("logistic") or {}
        out.append({"offer_id": pr.get("offer_id") or "", "sku": pr.get("sku"), "name": pr.get("name") or "",
                    "qty": pr.get("quantity") or 1, "reason": r.get("return_reason_name") or "—",
                    "type": r.get("type") or "", "schema": r.get("schema") or "", "status": st,
                    "posting": r.get("posting_number") or "",
                    "date": lg.get("return_date") or lg.get("final_moment") or r.get("created_at") or ""})
    by_offer = defaultdict(Counter)
    for x in out:
        by_offer[x["offer_id"]][x["reason"]] += int(x["qty"] or 1)
    res = {"from": frm.date().isoformat(), "to": to.date().isoformat(), "filter": used, "errors": errors,
           "count": len(out), "by_offer": {k: dict(v) for k, v in by_offer.items()}, "returns": out,
           "raw_sample": rows[:3]}
    with open(RETURNS_PATH, "w", encoding="utf-8") as fh:
        json.dump(res, fh, ensure_ascii=False, indent=1)
    return res


def cmd_ozon_returns() -> int:
    try:
        r = fetch_returns()
    except Exception as exc:  # noqa: BLE001
        print(f"ozon-returns: ОШИБКА — {exc}")
        return 1
    print(f"Возвраты Ozon {r['from']}..{r['to']}: {r['count']} (фильтр {r['filter']})")
    for e in r["errors"]:
        print("  не сработал", e)
    for oid, reasons in sorted(r["by_offer"].items(), key=lambda x: -sum(x[1].values())):
        print(f"  {oid}: " + "; ".join(f"{k} ×{v}" for k, v in reasons.items()))
    return 0


# ---------------------------------------------------------------- юнит-экономика
def _master() -> dict:
    import openpyxl
    import price_audit
    ws = openpyxl.load_workbook(price_audit.MASTER, read_only=True, data_only=True)["Товары"]
    rows = list(ws.iter_rows(values_only=True))
    h = [str(c or "").strip() for c in rows[0]]
    ic = h.index(price_audit.COST_HEADER) if price_audit.COST_HEADER in h else None
    ip, iname = h.index("Цена, ₽"), h.index("Название товара")
    out = {}
    for r in rows[1:]:
        if not r[0]:
            continue
        out[str(r[0]).strip().lower()] = {"art": str(r[0]).strip(), "name": r[iname] or "",
                                          "price": price_audit._num(r[ip]),
                                          "cost": price_audit._num(r[ic]) if ic is not None else None}
    return out


def build(fin: dict, rets: dict) -> list:
    import price_audit
    master = _master()
    try:
        ign = {k.lower() for k in json.load(open(os.path.join(DATA, "audit_ignore.json"), encoding="utf-8"))
               if not k.startswith("_")}
    except Exception:  # noqa: BLE001
        ign = set()
    fees = {k.lower(): v for k, v in (fin.get("fees_by_type") or {}).items()}
    reasons = {k.lower(): v for k, v in (rets.get("by_offer") or {}).items()}
    rows = []
    for oid, it in (fin.get("items") or {}).items():
        k = oid.lower()
        if k in ign or k.startswith("sku:"):
            continue
        m = master.get(k, {})
        cost = m.get("cost")
        units = max(int(it.get("sold", 0)) - int(it.get("returned", 0)), 0)
        rev = max(float(it.get("revenue") or 0), 0.0)
        tax = rev * price_audit.TAX
        profit = float(it.get("net") or 0) - tax - (cost or 0) * units
        f = fees.get(k, {})
        paid = sum(v for n, v in f.items() if n in PAID_PROGRAMS)
        rows.append({"offer_id": oid, "name": m.get("name", ""), "price": m.get("price"), "cost": cost,
                     "sold": it.get("sold", 0), "returned": it.get("returned", 0), "revenue": rev,
                     "commission": it.get("commission", 0), "logistics": it.get("logistics", 0),
                     "services": it.get("services", 0), "paid_programs": round(paid, 2),
                     "tax": round(tax, 2), "profit": round(profit, 2), "no_cost": cost is None,
                     "fees": f, "reasons": reasons.get(k, {})})
    rows.sort(key=lambda r: r["profit"])
    return rows


def _write_xlsx(rows: list, period: str) -> None:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Товары за 30 дней"
    head = ["Артикул", "Название", "Цена Ozon", "Закупка", "Продано", "Возвраты", "Выручка", "Комиссия", "Логистика",
            "Услуги всего", "из них платные программы", "Налог 6%", "Прибыль за месяц", "Что съело (услуги)", "Причины возвратов"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
    red = PatternFill("solid", fgColor="F4CCCC")
    for r in rows:
        top = "; ".join(f"{n} {v:,.0f}".replace(",", " ") for n, v in sorted(r["fees"].items(), key=lambda x: x[1])[:4] if v < 0)
        why = "; ".join(f"{n} ×{v}" for n, v in r["reasons"].items())
        ws.append([r["offer_id"], r["name"], r["price"], r["cost"] if not r["no_cost"] else "нет закупки", r["sold"],
                   r["returned"], r["revenue"], r["commission"], r["logistics"], r["services"], r["paid_programs"],
                   r["tax"], r["profit"], top, why])
        if r["profit"] < 0:
            for c in ws[ws.max_row]:
                c.fill = red
    for col, w in zip("ABCDEFGHIJKLMNO", (18, 46, 10, 10, 8, 8, 10, 10, 10, 10, 12, 9, 12, 60, 50)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C2"
    ws2 = wb.create_sheet("Как считается")
    for line in (f"Период: {period} (30 дней, фактические начисления Ozon /v1/finance/accrual/by-day).",
                 "Прибыль = начисления Ozon по товару (выручка − комиссия − логистика − услуги) − налог 6% с выручки "
                 "− закупка × (продано − возвращено).",
                 "Почему месяц, а не неделя: расходы Ozon списывает сразу, а выручку — при выкупе, поэтому недельный минус почти всегда ложный.",
                 "Платные программы: " + ", ".join(PAID_PROGRAMS) + ".",
                 "Закупка — колонка «Закупка, ₽» в data/master_control.xlsx; где её нет, прибыль без закупки (завышена)."):
        ws2.append([line])
    ws2.column_dimensions["A"].width = 140
    wb.save(REPORT_PATH)


def _fmt(v) -> str:
    return f"{v:,.0f}".replace(",", " ")


def message(rows: list, period: str) -> str:
    loss = [r for r in rows if r["profit"] < 0 and (r["sold"] or r["returned"])]
    paid = [r for r in rows if r["revenue"] > 0 and (r["price"] or 0) < CHEAP_PRICE
            and -r["paid_programs"] > PAID_SHARE_ALERT * r["revenue"]]
    lines = [f"📉 Ozon: юнит-экономика за 30 дней ({period})"]
    if not loss:
        lines.append("Убыточных товаров за месяц нет ✅")
    else:
        lines.append(f"В минусе за месяц: {len(loss)}")
        for r in loss[:12]:
            top = sorted(((n, v) for n, v in r["fees"].items() if v < 0), key=lambda x: x[1])[:2]
            ate = ", ".join(f"{n} {_fmt(-v)} ₽" for n, v in top)
            s = (f"• {r['offer_id']} — {_fmt(r['profit'])} ₽ (продано {r['sold']}, возвратов {r['returned']}"
                 + (", закупки нет в мастере" if r["no_cost"] else "") + ")")
            if r["returned"] > r["sold"]:
                s += "\n  ⚠️ возврат продажи прошлого месяца — минус может быть ложным"
            if ate:
                s += f"\n  съели: {ate}"
            if r["reasons"]:
                s += "\n  возвраты: " + "; ".join(f"{n} ×{v}" for n, v in r["reasons"].items())
            lines.append(s)
        if len(loss) > 12:
            lines.append(f"… и ещё {len(loss) - 12} — в data/ozon_unit_report.xlsx")
    if paid:
        lines.append(f"\n💸 Дешёвые товары, где «Звёздные товары»/сбор отзывов съедают >{PAID_SHARE_ALERT:.0%} выручки:")
        for r in paid[:10]:
            lines.append(f"• {r['offer_id']}: {_fmt(-r['paid_programs'])} ₽ из {_fmt(r['revenue'])} ₽ выручки")
        lines.append("Их можно отключить для этих товаров в кабинете Ozon (Продвижение).")
    lines.append("\nПолная таблица: data/ozon_unit_report.xlsx. Напишите Claude, если что-то разобрать.")
    return "\n".join(lines)


def cmd_ozon_unit_weekly(send: bool = True) -> int:
    import price_audit
    try:
        fin = price_audit.fetch_ozon_finance()
    except Exception as exc:  # noqa: BLE001
        print(f"ozon-unit-weekly: начисления Ozon не получены — {exc}")
        return 1
    try:
        rets = fetch_returns()
    except Exception as exc:  # noqa: BLE001
        print(f"возвраты не получены (продолжаю без причин): {exc}")
        rets = {}
    rows = build(fin, rets)
    period = f"{fin['from']}…{fin['to']}"
    _write_xlsx(rows, period)
    text = message(rows, period)
    print(text)
    if send and os.getenv("GITHUB_ACTIONS"):
        import direct_tools
        direct_tools._tg_owner(text)
    with open(STATE_PATH, "w", encoding="utf-8") as fh:
        json.dump({"sent": datetime.now(timezone.utc).isoformat(),
                   "loss": [r["offer_id"] for r in rows if r["profit"] < 0]}, fh, ensure_ascii=False)
    return 0
