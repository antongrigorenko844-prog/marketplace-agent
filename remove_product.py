"""
Правка data/master_control.xlsx для команды remove-product ("снять товар
отовсюду"): ставит "Остаток, шт." = 0 у указанных артикулов.

Остаток 0 в пульте — единый "рубильник": после sync-master-control он
попадает в рабочие каталоги, push-stock/push-wb-stock обнуляют витрины Ozon и
WB, а build-avito-catalog не включает такие товары в фид (объявление
закрывается). Сам товар из файлов НЕ удаляется — вернуть в продажу = вписать
остаток обратно и запустить update-cards/push-stock.
"""
import os
from typing import List, Tuple

from openpyxl import load_workbook

MASTER_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data", "master_control.xlsx")
COL_OFFER_ID = 1
COL_QUANTITY = 10  # "Остаток, шт." — как в sync_master_control.py
MAX_ARTICLES = 20  # защита от случайного "снять всё"


def zero_stock(offer_ids: List[str], dry_run: bool, path: str = MASTER_PATH) -> Tuple[List[Tuple[str, object]], List[str]]:
    """
    Возвращает ([(артикул, прежний остаток)], [не найденные артикулы]).
    При dry_run файл не сохраняется.
    """
    wb = load_workbook(path)
    ws = wb["Товары"]
    wanted = {o: o for o in offer_ids}
    found: List[Tuple[str, object]] = []
    for row in ws.iter_rows(min_row=2):
        cell = row[COL_OFFER_ID - 1]
        if cell.value is None:
            continue
        offer = str(cell.value).strip()
        if offer in wanted:
            qty_cell = row[COL_QUANTITY - 1]
            found.append((offer, qty_cell.value))
            qty_cell.value = 0
    missing = [o for o in offer_ids if o not in {f[0] for f in found}]
    if found and not dry_run:
        wb.save(path)
    return found, missing
