"""
Проверка data/master_control.xlsx ПЕРЕД отправкой на маркетплейсы
(validate-master; автоматически вызывается в начале update-cards).

Ловит типичные ошибки, которые иначе вскрываются уже после отправки:
пустая/нулевая цена, пустое название, нечисловой или отрицательный остаток,
название WB длиннее 60 символов, цена до скидки меньше цены, дубли артикулов,
отличающиеся только регистром/пробелами (на Mac-папках photos/ это ломает
фото, см. FOLDER_ALIASES в catalog_editor.py), товар без фото.

ОШИБКА (error) — то, что почти наверняка сломает отправку/покажет на витрине
неверное. ПРЕДУПРЕЖДЕНИЕ (warn) — стоит посмотреть, но не обязательно
ошибка. Тестовые и снятые с продажи строки (остаток 0 при пустых статусах
или "Тестовый товар") пропускаются. Файл НИКОГДА не меняется — только чтение.
"""
import os
from typing import Dict, List, Tuple

from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.abspath(__file__))
MASTER_PATH = os.path.join(ROOT, "data", "master_control.xlsx")

# Те же номера колонок, что в sync_master_control.py (1-based).
COL_OFFER_ID, COL_STATUS_OZON, COL_STATUS_WB, COL_FILES_COUNT = 1, 2, 3, 5
COL_NAME, COL_NAME_WB, COL_PRICE, COL_OLD_PRICE, COL_QUANTITY, COL_PRICE_WB = 6, 7, 8, 9, 10, 11
COL_DESCRIPTION, COL_NOTES = 12, 14

WB_NAME_MAX = 60

# Служебные строки (нужны для привязки/редиректа карточек, сами не продаются) — не проверяем.
SERVICE_ARTICLES = {"0am927769d-plata", "testwb001"}
COL_FLAG_OZON, COL_FLAG_WB = 17, 18  # «Ozon (да/нет)», «WB (да/нет)»


def _num(value):
    """Число или None. Пустое -> None; нечисловое -> ValueError."""
    if value is None or str(value).strip() == "":
        return None
    return float(str(value).replace(" ", "").replace("\xa0", "").replace(",", "."))


def _norm_article(a: str) -> str:
    return "".join(str(a).split()).lower().replace("‑", "-")


def validate(path: str = MASTER_PATH) -> Tuple[List[str], List[str]]:
    """Возвращает (errors, warnings) — списки готовых строк для вывода."""
    errors: List[str] = []
    warns: List[str] = []
    if not os.path.exists(path):
        return [f"Нет файла {path}"], []

    ws = load_workbook(path, read_only=True, data_only=True)["Товары"]
    seen: Dict[str, str] = {}
    checked = 0

    for r in ws.iter_rows(min_row=2, values_only=True):
        r = list(r) + [None] * 16
        offer = r[COL_OFFER_ID - 1]
        if offer is None or str(offer).strip() == "":
            continue
        offer = str(offer).strip()
        st_oz = str(r[COL_STATUS_OZON - 1] or "").strip()
        st_wb = str(r[COL_STATUS_WB - 1] or "").strip()
        notes = str(r[COL_NOTES - 1] or "")
        name_wb_cell = str(r[COL_NAME_WB - 1] or "")

        if "Тестов" in notes or "Тестов" in name_wb_cell:
            continue  # TEST001 / TESTWB001 не публикуются
        if offer.lower() in SERVICE_ARTICLES:
            continue  # служебные строки

        key = _norm_article(offer)
        if key in seen and seen[key] != offer:
            warns.append(f"{offer}: похож на «{seen[key]}» (отличается только регистром/пробелами) — убедитесь, что это два разных товара")
        seen.setdefault(key, offer)

        # Канал, отключённый в мастере («нет»), не проверяем: на нём товар не продаём.
        off_oz = str(r[COL_FLAG_OZON - 1] or "").strip().lower() == "нет"
        off_wb = str(r[COL_FLAG_WB - 1] or "").strip().lower() == "нет"
        on_ozon = st_oz != "" and not off_oz
        on_wb = st_wb != "" and not off_wb
        checked += 1

        # --- остаток
        qty = None
        try:
            qty = _num(r[COL_QUANTITY - 1])
        except ValueError:
            errors.append(f"{offer}: остаток «{r[COL_QUANTITY - 1]}» — не число")
        if qty is not None and qty < 0:
            errors.append(f"{offer}: остаток отрицательный ({qty:g})")
        if qty is None and (on_ozon or on_wb):
            warns.append(f"{offer}: остаток не задан — на Ozon/WB/Avito не обновится")

        sold_out = qty == 0  # снят/закончился — цену не требуем

        # --- цена Ozon / сайт / Avito
        try:
            price = _num(r[COL_PRICE - 1])
        except ValueError:
            price = None
            errors.append(f"{offer}: цена «{r[COL_PRICE - 1]}» — не число")
        else:
            if on_ozon and not sold_out and (price is None or price <= 0):
                errors.append(f"{offer}: нет цены (Ozon/сайт/Avito)")
        try:
            old = _num(r[COL_OLD_PRICE - 1])
            if price and old and old < price:
                warns.append(f"{offer}: цена до скидки ({old:g}) меньше цены ({price:g})")
        except ValueError:
            errors.append(f"{offer}: цена до скидки «{r[COL_OLD_PRICE - 1]}» — не число")

        # --- цена WB
        if on_wb and not sold_out:
            try:
                pwb = _num(r[COL_PRICE_WB - 1])
                if pwb is None or pwb <= 0:
                    errors.append(f"{offer}: нет цены WB")
            except ValueError:
                errors.append(f"{offer}: цена WB «{r[COL_PRICE_WB - 1]}» — не число")

        # --- названия
        name = str(r[COL_NAME - 1] or "").strip()
        if on_ozon and not name:
            errors.append(f"{offer}: пустое название")
        if on_wb and len(name_wb_cell.strip()) > WB_NAME_MAX:
            warns.append(f"{offer}: название WB {len(name_wb_cell.strip())} симв. (лимит {WB_NAME_MAX})")

        # --- описание, фото
        if (on_ozon or on_wb) and not str(r[COL_DESCRIPTION - 1] or "").strip():
            warns.append(f"{offer}: пустое описание")
        try:
            files = _num(r[COL_FILES_COUNT - 1])
        except ValueError:
            files = None
        if (on_ozon or on_wb) and not sold_out and (files is None or files <= 0):
            warns.append(f"{offer}: нет фото")

    if checked == 0:
        errors.append("В master_control.xlsx не найдено ни одной строки с артикулом")
    _check_dims(path, warns)
    return errors, warns


def _check_dims(path: str, warns: List[str]) -> None:
    """Вес/габариты: заполнять все четыре колонки сразу и только в мастере
    (не в data/ozon_overrides.json) — иначе Ozon и WB разойдутся."""
    import json

    try:
        import master_dims

        rows = master_dims.read_rows(path)
        full = master_dims.load(path)
    except Exception:
        return
    for art, vals in rows.items():
        if art not in full:
            warns.append(
                f"{art}: вес/габариты заполнены не полностью (нужны все четыре: «Вес, кг», «Длина, см», "
                f"«Ширина, см», «Высота, см», больше нуля) — пока не заполнены, на Ozon/WB не уйдут"
            )
    try:
        with open(os.path.join(ROOT, "data", "ozon_overrides.json"), encoding="utf-8") as f:
            overrides = json.load(f)
    except (OSError, ValueError):
        return
    full_low = {a.lower() for a in full}
    for art, ov in overrides.items():
        if isinstance(ov, dict) and ov.get("dims_mm") and art.lower() not in full_low:
            warns.append(
                f"{art}: габариты заданы только в data/ozon_overrides.json — уйдут на Ozon, но не на WB; "
                f"перенесите их в колонки «Вес, кг», «Длина/Ширина/Высота, см» мастера"
            )


def run(path: str = MASTER_PATH, quiet_ok: bool = False) -> int:
    errors, warns = validate(path)
    for e in errors:
        print(f"  ОШИБКА: {e}")
    for w in warns:
        print(f"  внимание: {w}")
    if errors:
        print(f"validate-master: ошибок {len(errors)}, предупреждений {len(warns)}.")
        return 1
    if warns:
        print(f"validate-master: критичных ошибок нет, предупреждений {len(warns)}.")
    elif not quiet_ok:
        print("validate-master: всё в порядке.")
    return 0


if __name__ == "__main__":
    raise SystemExit(run())
