"""
Контент-рейтинг карточек Ozon (ТОЛЬКО ЧТЕНИЕ — в кабинете ничего не меняется).

Ozon по каждой карточке считает контент-рейтинг 0–100 и говорит, чего не хватает:
медиа (фото, видео, видеообложка), атрибуты, описание/rich-контент и т.п.
Метод: POST /v1/product/rating-by-sku (до 100 SKU за запрос).

Итог: data/ozon_content_rating.xlsx
  «Карточки»            — каждая карточка: рейтинг, оценки по группам, что не выполнено,
                          какие атрибуты заполнить; порядок — по выгоде (продажи × цена).
  «Чего не хватает»     — какие условия не выполнены чаще всего (что чинить массово).
  «Как читать»          — пояснения.
Сырой ответ первой пачки — data/ozon_content_rating_raw.json (для сверки разбора).

Запуск: data/run_request.txt = ozon-content-rating (или Actions -> tools).
"""
from __future__ import annotations

import json
import os
from collections import Counter
from typing import Dict, List

DATA = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data")


def _master() -> Dict[str, dict]:
    import openpyxl

    ws = openpyxl.load_workbook(os.path.join(DATA, "master_control.xlsx"), read_only=True, data_only=True)["Товары"]
    rows = list(ws.iter_rows(values_only=True))
    head = [str(h or "") for h in rows[0]]
    out = {}
    for r in rows[1:]:
        if not r[0]:
            continue
        d = dict(zip(head, r))
        out[str(r[0]).strip()] = {
            "name": d.get("Название товара"),
            "price": d.get("Цена, ₽") or 0,
            "stock": d.get("Остаток, шт.") or 0,
        }
    return out


def _sales() -> Dict[str, int]:
    try:
        items = json.load(open(os.path.join(DATA, "ozon_finance.json"), encoding="utf-8")).get("items") or {}
    except Exception:  # noqa: BLE001
        return {}
    return {k: int(v.get("sold") or 0) - int(v.get("returned") or 0) for k, v in items.items()}


def _skus(info: dict) -> List[int]:
    """SKU карточки: в /v3/product/info/list — поле sku или sources[].sku."""
    out = []
    if info.get("sku"):
        out.append(int(info["sku"]))
    for s in info.get("sources") or []:
        if s.get("sku"):
            out.append(int(s["sku"]))
    for key in ("fbo_sku", "fbs_sku"):
        if info.get(key):
            out.append(int(info[key]))
    return list(dict.fromkeys(out))


def cmd_ozon_content_rating() -> int:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    import ozon_client

    master = _master()
    sales = _sales()
    plist = ozon_client.list_products()
    offer_ids = [str(p.get("offer_id")) for p in plist if p.get("offer_id")]
    infos = ozon_client.get_product_info_list(offer_ids)
    by_sku: Dict[int, dict] = {}
    for inf in infos:
        if inf.get("_error"):
            print("Ошибка info/list:", inf["_error"][:300])
            continue
        oid = str(inf.get("offer_id") or "")
        for s in _skus(inf):
            by_sku.setdefault(s, {"offer_id": oid, "images": len(inf.get("images") or []),
                                  "name": inf.get("name")})
    skus = list(by_sku)
    print(f"Карточек Ozon: {len(offer_ids)}, SKU для рейтинга: {len(skus)}")

    results: Dict[str, dict] = {}
    raw_saved = False
    for i in range(0, len(skus), 100):
        batch = skus[i: i + 100]
        data = ozon_client._post("/v1/product/rating-by-sku", {"skus": batch})
        if not raw_saved:
            json.dump(data, open(os.path.join(DATA, "ozon_content_rating_raw.json"), "w", encoding="utf-8"),
                      ensure_ascii=False, indent=1)
            raw_saved = True
        for p in data.get("products") or []:
            sku = int(p.get("sku") or 0)
            meta = by_sku.get(sku)
            if not meta:
                continue
            oid = meta["offer_id"]
            # у товара может быть несколько SKU (FBO/FBS) — берём первый, рейтинг у них общий
            results.setdefault(oid, {"sku": sku, "rating": p.get("rating"), "groups": p.get("groups") or [],
                                     "images": meta["images"]})

    fail_counter: Counter = Counter()
    attr_counter: Counter = Counter()
    group_names: List[str] = []
    rows = []
    for oid, r in results.items():
        m = master.get(oid, {})
        sold = max(0, sales.get(oid, 0))
        price = m.get("price") or 0
        grp = {}
        todo, attrs = [], []
        for g in r["groups"]:
            gname = g.get("name") or g.get("key") or "?"
            if gname not in group_names:
                group_names.append(gname)
            grp[gname] = g.get("rating")
            for c in g.get("conditions") or []:
                if not c.get("fulfilled"):
                    desc = (c.get("description") or c.get("key") or "").strip()
                    todo.append(f"{gname}: {desc}" + (f" (+{c.get('cost'):g})" if isinstance(c.get("cost"), (int, float)) else ""))
                    fail_counter[f"{gname}: {desc}"] += 1
            for a in g.get("improve_attributes") or []:
                an = a.get("name") or str(a.get("id"))
                attrs.append(an)
                attr_counter[an] += 1
        rows.append({
            "oid": oid, "name": m.get("name") or by_sku.get(r["sku"], {}).get("name"), "stock": m.get("stock"),
            "sold": sold, "price": price, "value": sold * price, "rating": r["rating"], "images": r["images"],
            "grp": grp, "todo": todo, "attrs": attrs,
        })
    rows.sort(key=lambda x: (-(x["value"] or 0), x["rating"] if x["rating"] is not None else 999))

    wb = openpyxl.Workbook()
    hdr_font, hdr_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="305496")

    def style(ws, widths):
        for c in ws[1]:
            c.font, c.fill = hdr_font, hdr_fill
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
        ws.freeze_panes = "C2"

    ws = wb.active
    ws.title = "Карточки"
    ws.append(["Артикул", "Название", "Рейтинг", "Остаток", "Продаж за мес", "Цена", "Фото, шт"]
              + group_names + ["Что не выполнено", "Атрибуты заполнить"])
    for x in rows:
        ws.append([x["oid"], x["name"], x["rating"], x["stock"], x["sold"], x["price"], x["images"]]
                  + [x["grp"].get(g) for g in group_names] + ["\n".join(x["todo"]), ", ".join(x["attrs"])])
    for row in ws.iter_rows(min_row=2):
        row[-2].alignment = row[-1].alignment = Alignment(wrap_text=True, vertical="top")
        r = row[2].value
        if isinstance(r, (int, float)):
            row[2].fill = PatternFill("solid", fgColor="F8CBAD" if r < 60 else "FFE699" if r < 80 else "C6EFCE")
    style(ws, [22, 45, 9, 9, 9, 10, 8] + [12] * len(group_names) + [70, 40])

    ws = wb.create_sheet("Чего не хватает")
    ws.append(["Условие", "У скольких карточек не выполнено"])
    for k, v in fail_counter.most_common():
        ws.append([k, v])
    ws.append([])
    ws.append(["Атрибут", "У скольких карточек не заполнен"])
    for k, v in attr_counter.most_common():
        ws.append([k, v])
    style(ws, [90, 18])

    ws = wb.create_sheet("Как читать")
    for t in [
        "Контент-рейтинг Ozon 0–100: красный < 60, жёлтый 60–79, зелёный ≥ 80. Чем выше, тем лучше карточка ранжируется и конвертирует.",
        "Порядок строк — по выгоде: продажи за месяц × цена. Сверху то, что приносит больше денег — улучшать в первую очередь.",
        "«Что не выполнено» — условия Ozon по группам (медиа, атрибуты, описание…), в скобках — сколько баллов даст выполнение.",
        "Ничего в кабинете не менялось — это только чтение.",
    ]:
        ws.append([t])
    ws.column_dimensions["A"].width = 140

    path = os.path.join(DATA, "ozon_content_rating.xlsx")
    wb.save(path)
    rated = [x["rating"] for x in rows if isinstance(x["rating"], (int, float))]
    avg = sum(rated) / len(rated) if rated else 0
    print(f"Рейтинг получен для {len(rows)} карточек, средний {avg:.1f}; ниже 60: {sum(1 for r in rated if r < 60)}")
    print("Чаще всего не выполнено:")
    for k, v in fail_counter.most_common(10):
        print(f"  {v:>3} × {k}")
    print(f"Сохранено: {path}")
    return 0
