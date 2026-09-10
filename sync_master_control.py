"""
sync_master_control.py — переносит правки из ЕДИНОГО файла-пульта
data/master_control.xlsx обратно в рабочие файлы конвейера:
  - data/ozon_catalog.xlsx        (строки со статусом "Действующий")
  - data/ozon_new_products.xlsx   (строки со статусом "Новый (черновик)")
а также, если в столбце "Фото" картинку заменили на другую, сохраняет
новую картинку как обложку в photos/<артикул>_1.<расширение>.

master_control.xlsx сам НЕ используется ни одним другим шагом конвейера —
он только человеко-читаемый обзор + место для правок. Реальные данные,
которые уходят в Ozon, всегда лежат в ozon_catalog.xlsx / ozon_new_products.xlsx
— эта функция и есть "мост" между ними: после правок в master_control.xlsx
запустите sync-master-control, а дальше конвейер работает как обычно
(attach-ozon-photos при смене обложки -> push-ozon-cards-dryrun -> push-ozon-cards,
и аналогично attach-ozon-new-photos -> push-ozon-new-cards-dryrun -> push-ozon-new-cards
для черновиков).

Столбцы master_control.xlsx (см. build_master_control.py), по позиции:
  1 Артикул (offer_id)
  2 Статус ("Действующий" / "Новый (черновик)")
  3 Фото (картинка, вставленная в ячейку)
  4 Кол-во файлов фото/видео (справочно, не читается обратно)
  5 Название товара
  6 Описание
  7 Хэштеги / Теги
  8 Заметки
  9 Файлы (папка photos/) (справочно, не читается обратно)

Пустая ячейка в столбцах название/описание/хэштеги/заметки означает
"не менять" — как и везде в этом проекте, обнулить значение так нельзя,
для явной очистки впишите один пробел.
"""
import hashlib
import logging
import os
import re

from openpyxl import load_workbook

logger = logging.getLogger("marketplace-agent.sync_master_control")

ROOT = os.path.dirname(os.path.abspath(__file__))
MASTER_PATH = os.path.join(ROOT, "data", "master_control.xlsx")
CATALOG_PATH = os.path.join(ROOT, "data", "ozon_catalog.xlsx")
NEW_PATH = os.path.join(ROOT, "data", "ozon_new_products.xlsx")
PHOTOS_DIR = os.path.join(ROOT, "photos")

STATUS_EXISTING = "Действующий"
STATUS_DRAFT = "Новый (черновик)"

# позиции столбцов в master_control.xlsx (1-based) — см. build_master_control.py
COL_OFFER_ID = 1
COL_STATUS = 2
COL_PHOTO = 3
COL_NAME = 5
COL_DESCRIPTION = 6
COL_HASHTAGS = 7
COL_NOTES = 8

IMAGE_EXTS = (".jpg", ".jpeg", ".png")


def _read_master():
    """
    Возвращает (rows, images):
      rows   — offer_id -> {"status","name","description","hashtags","notes"}
      images — offer_id -> raw bytes картинки, вставленной в столбец "Фото"
    """
    wb = load_workbook(MASTER_PATH)
    ws = wb["Товары"] if "Товары" in wb.sheetnames else wb.active

    row_to_offer = {}
    rows = {}
    for row in ws.iter_rows(min_row=2):
        offer_cell = row[COL_OFFER_ID - 1]
        if not offer_cell.value:
            continue
        offer_id = str(offer_cell.value).strip()
        row_idx = offer_cell.row
        row_to_offer[row_idx] = offer_id
        status = str(row[COL_STATUS - 1].value or "").strip()
        rows[offer_id] = {
            "status": status,
            "name": row[COL_NAME - 1].value,
            "description": row[COL_DESCRIPTION - 1].value,
            "hashtags": row[COL_HASHTAGS - 1].value,
            "notes": row[COL_NOTES - 1].value,
        }

    images = {}
    for img in getattr(ws, "_images", []):
        anchor = getattr(img, "anchor", None)
        row0 = None
        if anchor is not None and hasattr(anchor, "_from") and anchor._from is not None:
            row0 = anchor._from.row
        elif isinstance(anchor, str):
            m = re.match(r"[A-Za-z]+(\d+)", anchor)
            if m:
                row0 = int(m.group(1)) - 1
        if row0 is None:
            continue
        offer_id = row_to_offer.get(row0 + 1)
        if not offer_id:
            continue
        try:
            data = img._data()
        except Exception as exc:  # noqa: BLE001 — одна нечитаемая картинка не должна ронять весь sync
            logger.warning("Не удалось прочитать картинку в строке %s (%s): %s", row0 + 1, offer_id, exc)
            continue
        images[offer_id] = data

    return rows, images


def _guess_ext(data: bytes) -> str:
    if data[:8] == b"\x89PNG\r\n\x1a\n":
        return ".png"
    if data[:3] == b"\xff\xd8\xff":
        return ".jpg"
    return ".png"


def _match_candidates(fname: str, offer_ids, case_sensitive: bool):
    base, ext = os.path.splitext(fname)
    f_cmp = fname if case_sensitive else fname.lower()
    e_cmp = ext if case_sensitive else ext.lower()
    out = []
    for k in offer_ids:
        k_cmp = k if case_sensitive else k.lower()
        if f_cmp == (k_cmp + e_cmp) or f_cmp.startswith(k_cmp + "_") or f_cmp.startswith(k_cmp + "-"):
            out.append(k)
    return out


def _match_known(fname: str, offer_ids):
    exact = _match_candidates(fname, offer_ids, case_sensitive=True)
    if exact:
        return max(exact, key=len)
    loose = _match_candidates(fname, offer_ids, case_sensitive=False)
    if not loose:
        return None
    return max(loose, key=len)


def _thumb_source_photo(offer_id: str, all_offer_ids):
    """
    Файл, из которого build_master_control.py строит маленький эскиз в
    столбце "Фото" для этого артикула — нужен, чтобы пересобрать такой же
    эскиз и сравнить с картинкой, вставленной в ячейку (см. _make_thumb_bytes).
    Повторяет ТУ ЖЕ логику подбора, что и сам build-скрипт: артикулы и файлы
    перебираются в алфавитном порядке имени файла, совпадение — сначала
    точное по регистру, затем без учёта регистра, при нескольких кандидатах
    побеждает самый длинный префикс (чтобы "0am325025H" не перехватывал
    файлы артикула "0am325025HX").
    """
    if not os.path.isdir(PHOTOS_DIR):
        return None
    files = sorted(os.listdir(PHOTOS_DIR))
    media = [f for f in files if os.path.splitext(f)[1].lower() in IMAGE_EXTS]
    for f in media:
        if _match_known(f, all_offer_ids) == offer_id:
            return f
    return None


def _make_thumb_bytes(path: str):
    """
    Пересобирает эскиз ТОЧНО так же, как build_master_control.py (140x140,
    JPEG качество 80) — чтобы сравнить с картинкой, вставленной в ячейку.
    Нужно потому, что в master_control.xlsx хранится не оригинал фото, а
    его маленькая уменьшенная копия для просмотра; сравнивать напрямую
    оригинал с этой копией нельзя — они разного размера и формата и НИКОГДА
    не совпадут побайтово, даже если пользователь ничего не менял.
    """
    try:
        from PIL import Image as PILImage
        import io

        img = PILImage.open(path).convert("RGB")
        img.thumbnail((140, 140))
        buf = io.BytesIO()
        img.save(buf, "JPEG", quality=80)
        return buf.getvalue()
    except Exception as exc:  # noqa: BLE001
        logger.warning("Не удалось пересобрать эскиз %s для сравнения: %s", path, exc)
        return None


def _sync_photos(images: dict, all_offer_ids) -> list:
    """
    Для каждой картинки, вставленной в столбец "Фото" master_control.xlsx,
    сравнивает её с ЗАНОВО ПОСТРОЕННЫМ эскизом текущей обложки (а не с самой
    обложкой — см. _make_thumb_bytes). Совпал эскиз -> пользователь ничего
    не менял, пропускаем. Не совпал -> пользователь вставил другую картинку,
    сохраняем её как новую обложку <артикул>_1.<ext>. Старый файл обложки
    (если он назывался иначе) НЕ удаляется — на всякий случай, чтобы никогда
    не потерять фото молча; название выводится в консоль, чтобы вы проверили
    вручную, не остался ли он лишним.
    Возвращает список (offer_id, новый_файл, старый_файл_или_None).
    """
    os.makedirs(PHOTOS_DIR, exist_ok=True)
    written = []
    for offer_id, data in images.items():
        new_hash = hashlib.sha256(data).hexdigest()

        current = _thumb_source_photo(offer_id, all_offer_ids)
        if current:
            expected_thumb = _make_thumb_bytes(os.path.join(PHOTOS_DIR, current))
            if expected_thumb is not None and hashlib.sha256(expected_thumb).hexdigest() == new_hash:
                continue  # это тот же автосгенерированный эскиз — фото не меняли

        ext = _guess_ext(data)
        target_name = f"{offer_id}_1{ext}"
        with open(os.path.join(PHOTOS_DIR, target_name), "wb") as fh:
            fh.write(data)
        written.append((offer_id, target_name, current))
    return written


def _sync_catalog_xlsx(path: str, columns_module, updates: dict) -> int:
    """
    Точечно обновляет name/description/hashtags/notes по offer_id в уже
    существующем xlsx, не трогая остальные колонки (цену, ссылки на фото,
    остаток, ТН ВЭД и т.д.) — они остаются как были.
    """
    if not updates or not os.path.exists(path):
        return 0

    col_index = {key: i for i, (key, _) in enumerate(columns_module.COLUMNS, start=1)}
    wb = load_workbook(path)
    ws = wb.active

    changed = 0
    for row in ws.iter_rows(min_row=2):
        offer_cell = row[0]
        if not offer_cell.value:
            continue
        offer_id = str(offer_cell.value).strip()
        upd = updates.get(offer_id)
        if not upd:
            continue
        row_changed = False
        for key in ("name", "description", "hashtags", "notes"):
            col_idx = col_index.get(key)
            if not col_idx:
                continue
            new_val = upd.get(key)
            if new_val in (None, ""):
                continue  # пусто в master_control = "не менять"
            cell = row[col_idx - 1]
            if cell.value != new_val:
                cell.value = new_val
                row_changed = True
        if row_changed:
            changed += 1

    if changed:
        wb.save(path)
    return changed


def run() -> int:
    if not os.path.exists(MASTER_PATH):
        print(f"Нет файла {MASTER_PATH} — сначала создайте его (build_master_control.py) и загрузите в репозиторий.")
        return 1

    import catalog_editor
    import ozon_new_products

    rows, images = _read_master()

    existing_updates = {oid: r for oid, r in rows.items() if r["status"] == STATUS_EXISTING}
    draft_updates = {oid: r for oid, r in rows.items() if r["status"] == STATUS_DRAFT}

    changed_existing = _sync_catalog_xlsx(CATALOG_PATH, catalog_editor, existing_updates)
    changed_draft = _sync_catalog_xlsx(NEW_PATH, ozon_new_products, draft_updates)
    written_photos = _sync_photos(images, set(rows.keys()))

    print(f"master_control.xlsx -> ozon_catalog.xlsx: обновлено строк {changed_existing}")
    print(f"master_control.xlsx -> ozon_new_products.xlsx: обновлено строк {changed_draft}")
    if written_photos:
        print(f"Обложки обновлены/добавлены в photos/ ({len(written_photos)}):")
        for oid, fname, old_file in written_photos:
            note = ""
            if old_file and old_file != fname:
                note = f"  (проверьте, не остался ли лишним старый файл {old_file})"
            print(f"  {oid}: {fname}{note}")
    else:
        print("Изменённых обложек в столбце 'Фото' не найдено.")

    print(
        "\nГотово. Дальше как обычно: attach-ozon-photos (если менялись обложки существующих) "
        "-> push-ozon-cards-dryrun -> push-ozon-cards; для черновиков — "
        "attach-ozon-new-photos -> push-ozon-new-cards-dryrun -> push-ozon-new-cards."
    )
    return 0


if __name__ == "__main__":
    import sys
    sys.exit(run())
