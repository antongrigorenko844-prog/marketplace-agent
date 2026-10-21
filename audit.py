"""
Ревизия площадок (audit): сводит data/master_control.xlsx с тем, что реально
лежит на Ozon, WB, Авито, Яндекс Маркете и сайте (Tilda), и показывает ТОЛЬКО
расхождения. Ничего на площадках не меняет — только читает.

Источники (всё берётся из файлов data/, которые обновляют fetch-команды):
  Ozon     — ozon_products.json (+ ozon_stocks.json, если снят живой остаток)
  WB       — wb_cards.json (карточки + цены)
  Авито    — avito_live.json (живые объявления, если удалось получить),
             иначе avito_feed.xlsx (то, что МЫ отдаём Авито);
             avito_autoload.json — ошибки последней автозагрузки
  Маркет   — yandex_market_offers.xlsx
  Tilda    — tilda_offers0_1.xml (то, что мы отправили на сайт)

Какие товары где ДОЛЖНЫ быть — колонки «Ozon / WB / Авито / Маркет / Сайт (да/нет)»
в самом data/master_control.xlsx (справа, 17-21). Пусто = «да» (у тестовых — «нет»).
Хотите убрать товар с какой-то площадки — впишите «нет» в его строке, и ревизия
перестанет считать это пропуском. Мастер пересобирается при каждом запуске, но
значения этих колонок build-master-control переносит по артикулу.

Результат: data/audit_report.xlsx (листы «Сводка», «Проблемы», «Матрица») и
data/audit_summary.md (короткий текст для лога/уведомления).
"""
import json
import os
import re
import xml.etree.ElementTree as ET
from collections import Counter, OrderedDict
from typing import Dict, List, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(ROOT, "data")

CHANNELS = ["Ozon", "WB", "Авито", "Маркет", "Tilda"]
YES, NO = "да", "нет"


def _p(name: str) -> str:
    return os.path.join(DATA, name)


def key(article) -> str:
    """Артикулы на площадках различают регистр и дефисы (Ea888gen3 и EA888gen3 — разные товары),
    поэтому сравниваем точно, обрезая только пробелы по краям."""
    return str(article or "").strip()


def loose(article) -> str:
    """Артикул без регистра/пробелов/знаков — только чтобы заметить «тот же товар, другой регистр»."""
    return re.sub(r"[^0-9A-ZА-Я]", "", str(article or "").upper())


def _num(x) -> Optional[float]:
    try:
        return float(str(x).replace(" ", "").replace("\xa0", "").replace(",", "."))
    except (TypeError, ValueError):
        return None


def _norm_text(s) -> str:
    return re.sub(r"\s+", " ", str(s or "")).strip().lower()


def _is_test(article: str, name: str) -> bool:
    return str(article).upper().startswith("TEST") or "тестов" in _norm_text(name)


# ----------------------------------------------------------------------------
# Чтение источников
# ----------------------------------------------------------------------------

def _rows(path: str, sheet: Optional[str] = None) -> List[tuple]:
    if not os.path.exists(path):
        return []
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    return list(ws.iter_rows(values_only=True))


def _json(name: str):
    path = _p(name)
    if not os.path.exists(path):
        return None
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def load_master() -> "OrderedDict[str, dict]":
    rows = _rows(_p("master_control.xlsx"), "Товары")
    out: "OrderedDict[str, dict]" = OrderedDict()
    if not rows:
        return out
    head = [str(h or "") for h in rows[0]]

    def col(prefix: str) -> int:
        for i, h in enumerate(head):
            if h.startswith(prefix):
                return i
        return -1

    c = {
        "art": col("Артикул (offer_id)"),
        "st_oz": col("Статус Ozon"),
        "st_wb": col("Статус WB"),
        "name": col("Название товара"),
        "name_wb": col("Название WB"),
        "price": col("Цена, ₽"),
        "stock": col("Остаток"),
        "price_wb": col("Цена WB"),
        "desc": col("Описание"),
        "f_Ozon": col("Ozon (да/нет)"),
        "f_WB": col("WB (да/нет)"),
        "f_Авито": col("Авито (да/нет)"),
        "f_Маркет": col("Маркет (да/нет)"),
        "f_Tilda": col("Сайт (да/нет)"),
        "sample": col("Образец для WB"),
        "weight": col("Вес, кг"),
        "length": col("Длина, см"),
        "width": col("Ширина, см"),
        "height": col("Высота, см"),
    }
    for r in rows[1:]:
        if c["art"] < 0 or not r[c["art"]]:
            continue
        art = str(r[c["art"]]).strip()

        def g(k):
            return r[c[k]] if c[k] >= 0 and c[k] < len(r) else None

        out[art] = {
            "art": art,
            "st_oz": g("st_oz"),
            "st_wb": g("st_wb"),
            "name": g("name") or "",
            "name_wb": g("name_wb") or "",
            "price": _num(g("price")),
            "stock": _num(g("stock")) or 0,
            "price_wb": _num(g("price_wb")),
            "desc": g("desc") or "",
            "flags": {ch: str(g("f_" + ch) or "").strip().lower() for ch in CHANNELS},
            "sample": str(g("sample") or "").strip(),
            "weight": _num(g("weight")),
            "length": _num(g("length")),
            "width": _num(g("width")),
            "height": _num(g("height")),
        }
    return out


def load_ozon() -> Dict[str, dict]:
    d = _json("ozon_products.json")
    if not d:
        return {}
    out: Dict[str, dict] = {}
    for p in d.get("list", []):
        out[key(p.get("offer_id"))] = {"art": p.get("offer_id"), "archived": p.get("archived", False)}
    for p in d.get("details", []):
        k = key(p.get("offer_id"))
        if k in out:
            out[k]["price"] = _num((p.get("price") or {}).get("price"))
    for p in d.get("names", []):
        k = key(p.get("offer_id"))
        if k in out:
            out[k]["name"] = p.get("name", "")
    live = _json("ozon_stocks.json")
    if live:
        for k, v in live.items():
            if k in out:
                out[k]["stock"] = v
    return out


def load_wb() -> Dict[str, dict]:
    d = _json("wb_cards.json")
    if not d:
        return {}
    out: Dict[str, dict] = {}
    for c in d.get("cards", []):
        out[key(c.get("vendorCode"))] = {"art": c.get("vendorCode"), "name": c.get("title", "")}
    for p in d.get("prices", []):
        k = key(p.get("vendorCode"))
        if k in out and p.get("sizes"):
            out[k]["price"] = _num(p["sizes"][0].get("discountedPrice"))
    return out


def load_avito_feed() -> Dict[str, dict]:
    rows = _rows(_p("avito_feed.xlsx"), "Объявления")
    out: Dict[str, dict] = {}
    if len(rows) < 5:
        return out
    head = [str(h or "") for h in rows[1]]
    try:
        i_id = head.index("Уникальный идентификатор объявления")
        i_title = head.index("Название объявления")
        i_price = head.index("Цена")
    except ValueError:
        return out
    for r in rows[4:]:
        if r[i_id]:
            out[key(r[i_id])] = {"art": r[i_id], "title": r[i_title], "price": _num(r[i_price])}
    return out


def load_avito_live(feed: Dict[str, dict]) -> Optional[Dict[str, dict]]:
    """Живые объявления Авито, привязанные к артикулам по названию из фида."""
    d = _json("avito_live.json")
    if not d:
        return None
    by_title = {_norm_text(v.get("title")): k for k, v in feed.items()}
    # точное соответствие «объявление Авито -> наш артикул» даёт отчёт автозагрузки (ad_id = наш артикул)
    by_id = {}
    for r in (_json("avito_autoload.json") or {}).get("items", []):
        if r.get("avito_id") and r.get("ad_id"):
            by_id[int(r["avito_id"])] = key(r["ad_id"])
    out: Dict[str, dict] = {}
    unmatched = []
    for it in d.get("items", []):
        k = by_id.get(int(it.get("id") or 0)) or by_title.get(_norm_text(it.get("title")))
        if not k:
            unmatched.append(it)
            continue
        out[k] = {"status": it.get("status"), "price": _num(it.get("price")), "url": it.get("url"), "id": it.get("id")}
    out["__unmatched__"] = unmatched  # type: ignore[assignment]
    return out


def load_avito_autoload() -> Dict[str, List[str]]:
    d = _json("avito_autoload.json") or {}
    out: Dict[str, List[str]] = {}
    for it in d.get("items", []):
        # только настоящие ошибки (type=error): «автоподстановка телефона» и «вид запчасти определён
        # автоматически» — это предупреждения Авито, объявление при них публикуется
        msgs = [m for m in it.get("messages", []) if str(m.get("type", "")).lower() == "error"]
        if msgs:
            out[key(it.get("ad_id"))] = [re.sub(r"<[^>]+>", " ", f"{m.get('title') or ''} {m.get('description') or ''}") for m in msgs]
    return out


def load_market() -> Dict[str, dict]:
    out: Dict[str, dict] = {}
    for r in _rows(_p("yandex_market_offers.xlsx"))[1:]:
        if r and r[0] and str(r[2] or "").lower() != "да":
            out[key(r[0])] = {"art": r[0], "name": r[1] or "", "category": r[4]}
    return out


def load_tilda() -> Dict[str, dict]:
    path = _p("tilda_offers0_1.xml")
    out: Dict[str, dict] = {}
    if not os.path.exists(path):
        return out
    for e in ET.parse(path).iter("Предложение"):
        out[key(e.findtext("Артикул"))] = {
            "art": e.findtext("Артикул"),
            "qty": _num(e.findtext("Количество")) or 0,
            "price": _num(e.findtext("Цены/Цена/ЦенаЗаЕдиницу")),
        }
    return out


# ----------------------------------------------------------------------------
# Где товар ДОЛЖЕН быть (колонки мастера «… (да/нет)»)
# ----------------------------------------------------------------------------

def channel_flags(master: "OrderedDict[str, dict]") -> Dict[str, dict]:
    """{артикул: {канал: 'да'/'нет'}}. Пусто в мастере = «да», у тестовых строк = «нет»."""
    flags: Dict[str, dict] = {}
    for art, m in master.items():
        test = _is_test(art, m["name"])
        d = {}
        for ch in CHANNELS:
            v = m["flags"].get(ch, "")
            d[ch] = v if v in (YES, NO) else (NO if test else YES)
        flags[art] = d
    return flags


# ----------------------------------------------------------------------------
# Сравнение
# ----------------------------------------------------------------------------

def compute() -> dict:
    master = load_master()
    flags = channel_flags(master)
    oz, wb_, feed, mk, tl = load_ozon(), load_wb(), load_avito_feed(), load_market(), load_tilda()
    live = load_avito_live(feed)
    unmatched_live = (live or {}).pop("__unmatched__", []) if live else []
    autoload = load_avito_autoload()

    have_data = {
        "Ozon": bool(oz), "WB": bool(wb_), "Авито": bool(feed), "Маркет": bool(mk), "Tilda": bool(tl),
    }
    present = {
        "Ozon": set(oz), "WB": set(wb_), "Маркет": set(mk), "Tilda": set(tl),
        "Авито": set(live) if live else set(feed),
    }

    loose_present = {ch: {loose(x): x for x in present[ch]} for ch in CHANNELS}
    problems: List[dict] = []

    def add(sev, art, ch, kind, detail, action):
        problems.append({"sev": sev, "art": art, "ch": ch, "kind": kind, "detail": detail, "action": action})

    matrix = []
    for art, m in master.items():
        k = key(art)
        f = flags.get(art, {})
        stock = m["stock"]
        row = {"art": art, "name": m["name"], "stock": stock, "price": m["price"], "cells": {}}
        for ch in CHANNELS:
            want = f.get(ch, YES) == YES
            here = k in present[ch]
            case_only = (not here) and loose(k) in loose_present[ch]
            row["cells"][ch] = ("есть" if here else ("регистр" if case_only else "нет")) if want else ("лишний" if here else "—")
            if not have_data[ch]:
                continue
            if want and case_only:
                add("НИЗКО", art, ch, "артикул отличается регистром", f"в мастере «{art}», на площадке «{loose_present[ch][loose(k)]}»", "привести к одному написанию (или оставить, если так заведено исторически)")
                continue
            if want and not here and stock > 0:
                action = {
                    "Ozon": "создать карточку на Ozon (new_products)",
                    "WB": "создать карточку WB: команда wb-create-missing-dryrun, затем wb-create-missing (tools)",
                    "Авито": "нет объявления на Авито — смотрите ошибки автозагрузки; для двигателя нужен другой шаблон категории",
                    "Маркет": "проверьте синхронизацию Ozon→Маркет или создайте оффер",
                    "Tilda": "проверьте push-tilda / видимость товара на сайте",
                }[ch]
                add("ВАЖНО", art, ch, "нет на площадке", f"остаток {stock:g}", action)
            elif not want and here:
                add("СРЕДНЕ", art, ch, "лишний", "в мастере для этой площадки стоит «нет», но товар на ней есть", "убрать с площадки (remove-product) или поставить «да»")
        # цены и названия
        o = oz.get(k)
        if o:
            if o.get("price") and m["price"] and abs(o["price"] - m["price"]) > 1:
                add("ВАЖНО", art, "Ozon", "цена", f"Ozon {o['price']:g} ≠ мастер {m['price']:g}", "update-cards (push-photos/push-ozon-cards) или проверьте ручную акцию")
            if o.get("name") and _norm_text(o["name"]) != _norm_text(m["name"]):
                add("СРЕДНЕ", art, "Ozon", "название", f"Ozon: «{o['name'][:60]}» ≠ мастер: «{m['name'][:60]}»", "update-cards")
            if o.get("stock") is not None and abs(o["stock"] - stock) > 0:
                add("СРЕДНЕ", art, "Ozon", "остаток", f"Ozon {o['stock']:g} ≠ мастер {stock:g}", "push-stock")
        w = wb_.get(k)
        if w:
            if m["name_wb"] and _norm_text(w.get("name")) != _norm_text(m["name_wb"]):
                add("СРЕДНЕ", art, "WB", "название", f"WB: «{str(w.get('name'))[:50]}» ≠ мастер: «{m['name_wb'][:50]}»", "update-cards")
            if w.get("price") and m["price_wb"] and abs(w["price"] - m["price_wb"]) > max(2, 0.01 * m["price_wb"]):
                add("ВАЖНО", art, "WB", "цена", f"WB {w['price']:g} ≠ мастер {m['price_wb']:g}", "push-wb-price")
        # Правило: на трёх маркетплейсах (Ozon, WB, Маркет) цена покупателя должна совпадать.
        if stock > 0 and m["price"] and m["price_wb"] and k in wb_ and abs(m["price"] - m["price_wb"]) > 1:
            add("СРЕДНЕ", art, "WB", "цена WB ≠ цена Ozon", f"Ozon {m['price']:g}, WB {m['price_wb']:g} ({m['price_wb'] / m['price']:.0%} от Ozon)", "выровнять: вписать цену Ozon в колонку «Цена WB» мастера и запустить update-cards")
        mm = mk.get(k)
        if mm and mm.get("name") and _norm_text(mm["name"]) != _norm_text(m["name"]):
            add("НИЗКО", art, "Маркет", "название", f"Маркет: «{mm['name'][:50]}» ≠ мастер: «{m['name'][:50]}»", "поправить в кабинете Маркета (названия через API не синхронизируются)")
        # Две ценовые ступени: маркетплейсы (Ozon/WB/Маркет — с комиссией) и Авито + сайт (без комиссии,
        # цена ≈ половина маркетплейсовой; минимальную цену Авито не учитываем).
        fa, tt = feed.get(k), tl.get(k)
        if fa and fa.get("price") and tt and tt.get("price") and abs(fa["price"] - tt["price"]) > 1 and stock > 0:
            add("СРЕДНЕ", art, "Авито/сайт", "цены Авито и сайта разные", f"Авито {fa['price']:g} ≠ сайт {tt['price']:g}", "должны совпадать — проверьте avito_site_price_overrides.xlsx")
        if stock > 0 and m["price"] and fa and fa.get("price"):
            half = int(m["price"] / 2 + 0.5)
            if abs(fa["price"] - m["price"]) <= 1:
                add("СРЕДНЕ", art, "Авито/сайт", "нет цены без комиссии", f"на Авито цена как на маркетплейсах ({fa['price']:g}); по правилу ≈ {half}", f"добавить {art} = {half} в data/avito_site_price_overrides.xlsx")
            elif abs(fa["price"] - half) > max(5, 0.03 * half):
                add("НИЗКО", art, "Авито/сайт", "цена не равна половине", f"Авито/сайт {fa['price']:g}, половина маркетплейса {half}", "оставить, если это сознательная цена")
        if live and k in live and k in feed and live[k].get("price") and feed[k].get("price"):
            if abs(live[k]["price"] - feed[k]["price"]) > 1:
                add("СРЕДНЕ", art, "Авито", "цена", f"на Авито {live[k]['price']:g} ≠ в нашем файле {feed[k]['price']:g}", "дождаться следующей автозагрузки / проверить её отчёт")
        if k in autoload:
            add("ВАЖНО", art, "Авито", "ошибка автозагрузки", " | ".join(autoload[k])[:300], "исправить поле в шаблоне/описании")
        matrix.append(row)

    # лишнее на площадках, чего нет в мастере
    mk_keys = {key(a) for a in master}
    for ch, src in (("Ozon", oz), ("WB", wb_), ("Маркет", mk), ("Tilda", tl)):
        for k, v in src.items():
            if k not in mk_keys and loose(k) not in {loose(a) for a in master}:
                add("СРЕДНЕ", v.get("art", k), ch, "нет в мастере", "товар есть на площадке, но в master_control его нет", "завести в мастер или убрать с площадки")
    for it in unmatched_live[:50]:
        add("СРЕДНЕ", it.get("title", "?"), "Авито", "объявление не из фида", f"id {it.get('id')} статус {it.get('status')}", "закрыть, если старое/дубль")

    return {
        "master": master, "problems": problems, "matrix": matrix,
        "have_data": have_data, "avito_source": "живые объявления" if live else "наш файл-фид",
    }


# ----------------------------------------------------------------------------
# Отчёт
# ----------------------------------------------------------------------------

def write_report(res: dict) -> dict:
    problems, matrix = res["problems"], res["matrix"]
    order = {"ВАЖНО": 0, "СРЕДНЕ": 1, "НИЗКО": 2}
    problems.sort(key=lambda p: (order[p["sev"]], p["ch"], str(p["art"])))
    by_sev = Counter(p["sev"] for p in problems)
    by_kind = Counter((p["ch"], p["kind"]) for p in problems)

    wb = Workbook()
    s = wb.active
    s.title = "Сводка"
    s.append(["Показатель", "Значение"])
    s.append(["Товаров в мастере", len(res["master"])])
    s.append(["С остатком > 0", sum(1 for m in matrix if m["stock"] > 0)])
    s.append(["Проблем ВАЖНО / СРЕДНЕ / НИЗКО", f"{by_sev['ВАЖНО']} / {by_sev['СРЕДНЕ']} / {by_sev['НИЗКО']}"])
    s.append(["Данные Авито взяты из", res["avito_source"]])
    for ch, ok in res["have_data"].items():
        if not ok:
            s.append([f"НЕТ ДАННЫХ: {ch}", "файл выгрузки пуст или отсутствует — проверка этого канала пропущена"])
    s.append([])
    s.append(["Канал / тип проблемы", "Сколько"])
    for (ch, kind), n in sorted(by_kind.items(), key=lambda x: -x[1]):
        s.append([f"{ch}: {kind}", n])
    s.column_dimensions["A"].width = 44
    s.column_dimensions["B"].width = 70

    ps = wb.create_sheet("Проблемы")
    ps.append(["Важность", "Артикул", "Канал", "Что не так", "Подробности", "Что делать"])
    fills = {"ВАЖНО": "F8CBAD", "СРЕДНЕ": "FFF2CC", "НИЗКО": "E2EFDA"}
    for p in problems:
        ps.append([p["sev"], p["art"], p["ch"], p["kind"], p["detail"], p["action"]])
        ps.cell(row=ps.max_row, column=1).fill = PatternFill("solid", fgColor=fills[p["sev"]])
    for i, w in enumerate([10, 30, 9, 22, 60, 70], 1):
        ps.column_dimensions[get_column_letter(i)].width = w
    ps.freeze_panes = "A2"
    ps.auto_filter.ref = ps.dimensions

    ms = wb.create_sheet("Матрица")
    ms.append(["Артикул", "Название", "Остаток", "Цена"] + CHANNELS)
    red, grn, gry = (PatternFill("solid", fgColor=c) for c in ("F8CBAD", "E2EFDA", "EDEDED"))
    for m in matrix:
        ms.append([m["art"], m["name"], m["stock"], m["price"]] + [m["cells"][ch] for ch in CHANNELS])
        for j, ch in enumerate(CHANNELS):
            c = ms.cell(row=ms.max_row, column=5 + j)
            c.fill = grn if c.value == "есть" else (gry if c.value == "—" else (PatternFill("solid", fgColor="FFF2CC") if c.value == "регистр" else red))
            c.alignment = Alignment(horizontal="center")
    for c in list(ms[1]) + list(ps[1]) + list(s[1]):
        c.font = Font(bold=True)
    for i, w in enumerate([26, 50, 9, 10, 9, 9, 9, 9, 9], 1):
        ms.column_dimensions[get_column_letter(i)].width = w
    ms.freeze_panes = "C2"
    ms.auto_filter.ref = ms.dimensions
    wb.save(_p("audit_report.xlsx"))

    lines = [f"Ревизия площадок: проблем ВАЖНО {by_sev['ВАЖНО']}, СРЕДНЕ {by_sev['СРЕДНЕ']}, НИЗКО {by_sev['НИЗКО']} (Авито: {res['avito_source']})"]
    for (ch, kind), n in sorted(by_kind.items(), key=lambda x: -x[1])[:12]:
        lines.append(f"  - {ch}: {kind} — {n}")
    for ch, ok in res["have_data"].items():
        if not ok:
            lines.append(f"  ! нет данных по каналу {ch} — не проверялся")
    text = "\n".join(lines)
    with open(_p("audit_summary.md"), "w", encoding="utf-8") as f:
        f.write(text + "\n")
    return {"text": text, "important": by_sev["ВАЖНО"], "total": len(problems)}


def _track_new(res: dict) -> List[str]:
    """Какие ВАЖНЫЕ проблемы появились с прошлой ревизии (состояние — data/audit_state.json)."""
    cur = sorted({f"{p['ch']}|{p['kind']}|{p['art']}" for p in res["problems"] if p["sev"] == "ВАЖНО"})
    prev = set((_json("audit_state.json") or {}).get("important", []))
    with open(_p("audit_state.json"), "w", encoding="utf-8") as f:
        json.dump({"important": cur}, f, ensure_ascii=False, indent=0)
    return [c for c in cur if c not in prev]


def run() -> dict:
    res = compute()
    out = write_report(res)
    new = _track_new(res)
    out["new_important"] = new
    if new:
        extra = f"НОВЫХ важных проблем с прошлой ревизии: {len(new)}"
        out["text"] += "\n" + extra
        print(out["text"])
        for n in new[:30]:
            print("  + " + n.replace("|", " / "))
        with open(_p("audit_summary.md"), "a", encoding="utf-8") as f:
            f.write(extra + "\n" + "\n".join("  + " + n.replace("|", " / ") for n in new[:30]) + "\n")
    else:
        print(out["text"])
    print("Подробности: data/audit_report.xlsx (после запуска сделайте Pull origin).")
    return out
