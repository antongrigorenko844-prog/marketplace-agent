"""
Яндекс Директ: проверка файла импорта, проверка посадочных страниц и ежедневная сводка в Telegram.

Источник правды для кампаний — файлы data/direct/*.tsv (тот же формат, что загружается в Директ
через «Загрузить из файла»: колонки «Название группы», «Фраза (с минус-словами)», «Минус-фразы на группу»,
«Заголовок 1/2», «Текст 1», «Ссылка», «Адреса быстрых ссылок», «Минус-фразы на кампанию» …).

Команды (main.py):
  --direct-lint     проверить все data/direct/*.tsv: мусорные/широкие фразы (как «китайский» из одного слова),
                    чужие коробки (DQ250/DQ500/0B5…), ремонтный трафик, конфликты фраз с минус-словами,
                    длины заголовков/текстов, UTM. Ошибки → сообщение владельцу в Telegram, прогон красный.
  --direct-check-links  открыть все страницы из кампаний (ссылки и быстрые ссылки): ответ 200, нет заглушек
                    (wa.me/71234567890, t.me/username, «renew subscription»), есть якоря (#katalog …).
                    Проблемы → Telegram (одно сообщение в день, только если что-то изменилось).
  --direct-report   сводка Директа за вчера и 7 дней: расход с НДС, клики, цена клика, конверсии,
                    + самые дорогие поисковые запросы вчера (чтобы сразу видеть мусор и дописать минус-слова).
                    Нужен секрет YANDEX_DIRECT_TOKEN (OAuth-токен Директа); без него команда молча выходит.
"""
import csv
import glob
import hashlib
import json
import os
import re
import time
from datetime import date, datetime, timedelta
from typing import Dict, List, Optional, Tuple

import requests

_ROOT = os.path.dirname(os.path.abspath(__file__))
DIRECT_DIR = os.path.join(_ROOT, "data", "direct")
STATE_PATH = os.path.join(_ROOT, "data", "direct_state.json")
OWNER_PATH = os.path.join(_ROOT, "data", "owner_chat.json")
SITE_HOST = "leoparts-akpp.ru"
UA = {"User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/141.0 Safari/537.36"}

# Лимиты Директа (текстово-графическое объявление)
LIMITS = {"Заголовок 1": 56, "Заголовок 2": 30, "Текст 1": 81}
SITELINK_TITLE_MAX = 30
SITELINK_DESC_MAX = 60
WORD_MAX = 22

# Признаки «нашей» фразы: коробка/деталь. Фраза без них — слишком широкая.
ANCHORS = ("dq200", "dq 200", "dsg", "дсг", "0am", "0cw", "мехатрон", "гидроблок", "плата", "эбу", "блок")
# Чужие коробки и нецелевой интент — в фразе быть не должно
OFF_TARGET = ("dq250", "dq500", "dq381", "dq400", "0b5", "02e", "0dd", "dsg 6", "дсг 6", "dsg6", "дсг6",
              "ремонт", "замена", "заменить", "ремкомплект", "прошивк", "диагностик", "адаптаци", "своими руками")
# Заглушки шаблона Тильды, которые не должны оказаться на рекламируемой странице
PLACEHOLDERS = ("wa.me/71234567890", "wa.me/11234567890", "t.me/username", "vk.me/username",
                "renew subscription", "renew your subscription")
STOPWORDS = {"в", "на", "с", "со", "для", "по", "от", "до", "из", "к", "о", "об", "и", "а", "но", "или", "за",
             "без", "под", "над", "при", "у", "же", "ли", "не", "как", "это", "что", "такое"}


# ---------------------------------------------------------------- общие мелочи

def _tg_owner(text: str) -> bool:
    token = os.getenv("TELEGRAM_BOT_TOKEN", "")
    try:
        chat = json.load(open(OWNER_PATH, encoding="utf-8")).get("chat_id")
    except (OSError, ValueError):
        chat = None
    if not token or not chat:
        print("Telegram не настроен — сообщение только в лог:\n" + text)
        return False
    try:
        for i in range(0, len(text), 3900):
            requests.post(f"https://api.telegram.org/bot{token}/sendMessage",
                          data={"chat_id": chat, "text": text[i:i + 3900], "disable_web_page_preview": "true"}, timeout=30)
        return True
    except requests.RequestException as exc:
        print(f"Telegram: {exc}")
        return False


def _state() -> dict:
    try:
        return json.load(open(STATE_PATH, encoding="utf-8"))
    except (OSError, ValueError):
        return {}


def _save_state(st: dict) -> None:
    os.makedirs(os.path.dirname(STATE_PATH), exist_ok=True)
    with open(STATE_PATH, "w", encoding="utf-8") as f:
        json.dump(st, f, ensure_ascii=False, indent=2, sort_keys=True)


def read_tsv(path: str) -> List[Dict[str, str]]:
    with open(path, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f, delimiter="\t"))
    if not rows:
        return []
    head = [h.strip() for h in rows[0]]
    return [dict(zip(head, r + [""] * (len(head) - len(r)))) for r in rows[1:] if any(c.strip() for c in r)]


def _split_list(s: str) -> List[str]:
    return [x.strip() for x in re.split(r"[,\n]", s or "") if x.strip()]


def _phrase_core(phrase: str) -> Tuple[str, List[str]]:
    """'мехатроник dq200 -бу -китай' -> ('мехатроник dq200', ['бу','китай'])"""
    parts = re.split(r"\s+-(?=\S)", phrase.strip())
    return parts[0].strip(), [p.strip() for p in parts[1:]]


def _tokens(core: str) -> List[str]:
    core = re.sub(r'["\[\]()]', " ", core.lower())
    return [t for t in core.split() if t]


def _stem(w: str) -> str:
    w = w.lower().lstrip("+")
    if re.search(r"\d", w) or len(w) <= 4:
        return w
    return w[:max(4, len(w) - 2)]


def _word_hits(minus_word: str, token: str) -> bool:
    """Совпадает ли минус-слово с токеном фразы (Директ учитывает все словоформы, если нет «!»)."""
    m = minus_word.lower().strip()
    t = token.lower().lstrip("+!")
    if m.startswith("!"):
        return t == m[1:]
    m = m.lstrip("+")
    if re.search(r"\d", m) or len(m) <= 4:
        return t == m
    s = _stem(m)
    return t.startswith(s) and len(t) - len(s) <= 4


def _minus_hits(minus: str, toks: List[str]) -> bool:
    words = [w for w in minus.split() if w.lower().lstrip("!+") not in STOPWORDS or w.startswith(("!", "+"))]
    if not words:
        return False
    return all(any(_word_hits(w, t) for t in toks) for w in words)


# ---------------------------------------------------------------- проверка файла импорта

def lint_file(path: str) -> Tuple[List[str], List[str]]:
    errors: List[str] = []
    warns: List[str] = []
    rows = read_tsv(path)
    need = ["Название группы", "Фраза (с минус-словами)", "Заголовок 1", "Текст 1", "Ссылка"]
    if not rows:
        return [f"файл пустой"], []
    miss = [c for c in need if c not in rows[0]]
    if miss:
        return [f"нет колонок: {', '.join(miss)}"], []

    camp_minus: List[str] = []
    for r in rows:
        camp_minus += _split_list(r.get("Минус-фразы на кампанию", ""))
    for m in camp_minus:
        ws = m.split()
        if len(ws) == 1 and ws[0].lower() in STOPWORDS:
            errors.append(f"минус-слово кампании «{m}» — стоп-слово, Директ его не примет; пишите «!{m}»")

    seen: Dict[str, str] = {}
    group_minus: Dict[str, List[str]] = {}
    for r in rows:
        g = r["Название группы"].strip()
        if r.get("Минус-фразы на группу", "").strip():
            group_minus[g] = _split_list(r["Минус-фразы на группу"])

    for i, r in enumerate(rows, start=2):
        g = r["Название группы"].strip()
        phrase = r["Фраза (с минус-словами)"].strip()
        if not phrase:
            continue
        where = f"стр.{i} [{g}] «{phrase}»"
        core, own_minus = _phrase_core(phrase)
        toks = _tokens(core)
        low = " " + " ".join(toks) + " "
        key = " ".join(sorted(t.lstrip("+!") for t in toks if t.lstrip("+!") not in STOPWORDS))
        if key in seen:
            warns.append(f"{where}: дубль фразы «{seen[key]}»")
        seen[key] = core
        content = [t for t in toks if t.lstrip("+!") not in STOPWORDS]
        is_code = len(content) == 1 and re.search(r"\d", content[0]) and (
            re.search(r"[a-zа-я]", content[0]) or re.fullmatch(r"\d{5,6}", content[0].lstrip("+!")))  # 01315, 006079
        if len(content) <= 1 and not is_code:
            errors.append(f"{where}: фраза из одного слова — соберёт весь интернет (так «китайский» дал 9 млн показов в прогнозе)")
        elif not any(a in low for a in ANCHORS):
            if g.lower().find("модел") < 0 or "мехатрон" not in low:
                warns.append(f"{where}: нет признака DQ200/DSG7/мехатроника — проверьте, не слишком ли широко")
        for bad in OFF_TARGET:
            if bad.startswith("замен") and "замена" in g.lower():
                continue   # отдельная группа «Замена платы»: замена = покупка новой платы
            if (" " + bad) in low:
                errors.append(f"{where}: нецелевое «{bad}» (чужая коробка или ремонт/услуга)")
                break
        if re.search(r"(^|\s)без\s", " " + core.lower() + " ") and "+без" not in core.lower():
            warns.append(f"{where}: предлог «без» меняет смысл — зафиксируйте как «+без»")
        for scope, lst in (("группы", group_minus.get(g, [])), ("кампании", camp_minus)):
            for m in lst:
                if _minus_hits(m, toks):
                    errors.append(f"{where}: конфликт с минус-словом {scope} «{m}» — фраза не будет показываться")
                    break
        for m in own_minus:
            if _minus_hits(m, toks):
                errors.append(f"{where}: своё минус-слово «{m}» совпадает с фразой")

    for i, r in enumerate(rows, start=2):
        b = (r.get("Ставка") or "").replace(",", ".").strip()
        if b:
            try:
                v = float(b)
                if not 0.3 <= v <= 60:
                    errors.append(f"стр.{i} [{r['Название группы']}] ставка {b} ₽ вне 0,3–60 ₽")
            except ValueError:
                errors.append(f"стр.{i} [{r['Название группы']}] ставка «{b}» — не число")

    # объявления (первая строка каждой группы несёт текст объявления)
    done = set()
    for i, r in enumerate(rows, start=2):
        g = r["Название группы"].strip()
        if g in done or not r.get("Заголовок 1"):
            continue
        done.add(g)
        for col, lim in LIMITS.items():
            v = r.get(col, "")
            n = len(re.sub(r"[!,.;:\"]", "", v))
            if n > lim:
                errors.append(f"[{g}] {col}: {n} > {lim} симв.: «{v}»")
            for w in v.split():
                if len(w) > WORD_MAX:
                    errors.append(f"[{g}] {col}: слово длиннее {WORD_MAX}: «{w}»")
        url = r.get("Ссылка", "")
        if SITE_HOST not in url:
            errors.append(f"[{g}] ссылка не на {SITE_HOST}: {url[:80]}")
        if "utm_source=" not in url:
            warns.append(f"[{g}] в ссылке нет UTM-меток")
        titles = [x for x in r.get("Заголовки быстрых ссылок", "").split("||") if x]
        descs = [x for x in r.get("Описания быстрых ссылок", "").split("||") if x]
        addrs = [x for x in r.get("Адреса быстрых ссылок", "").split("||") if x]
        if titles and len(titles) != len(addrs):
            errors.append(f"[{g}] быстрых ссылок {len(titles)}, а адресов {len(addrs)}")
        if len(set(addrs)) != len(addrs):
            errors.append(f"[{g}] адреса быстрых ссылок повторяются")
        for t in titles:
            if len(t) > SITELINK_TITLE_MAX:
                errors.append(f"[{g}] заголовок быстрой ссылки > {SITELINK_TITLE_MAX}: «{t}»")
        for d in descs:
            if len(d) > SITELINK_DESC_MAX:
                errors.append(f"[{g}] описание быстрой ссылки > {SITELINK_DESC_MAX}: «{d}»")
        if sum(len(t) for t in titles[:4]) > 66:
            warns.append(f"[{g}] первые 4 быстрые ссылки длиннее 66 символов в сумме — часть не покажется")
    return errors, warns


def cmd_direct_lint(paths: Optional[List[str]] = None) -> int:
    paths = paths or sorted(glob.glob(os.path.join(DIRECT_DIR, "*.tsv")))
    if not paths:
        print(f"Нет файлов {DIRECT_DIR}/*.tsv")
        return 0
    bad = 0
    report = []
    for p in paths:
        errs, warns = lint_file(p)
        name = os.path.basename(p)
        print(f"\n=== {name}: ошибок {len(errs)}, предупреждений {len(warns)}")
        for e in errs:
            print("  ОШИБКА:", e)
        for w in warns:
            print("  внимание:", w)
        if errs:
            bad += 1
            report.append(f"❌ {name}: {len(errs)} ошиб.\n" + "\n".join("• " + e for e in errs[:15])
                          + (f"\n… и ещё {len(errs) - 15}" if len(errs) > 15 else ""))
    if report and os.getenv("GITHUB_ACTIONS"):
        _tg_owner("Директ: файл импорта не прошёл проверку — не загружайте его, пока не поправим.\n\n" + "\n\n".join(report))
    return 1 if bad else 0


# ---------------------------------------------------------------- проверка посадочных

def active_files() -> List[str]:
    """Файлы кампаний, которые сейчас загружены в Директ (data/direct/active.txt); иначе все."""
    try:
        names = [x.strip() for x in open(os.path.join(DIRECT_DIR, "active.txt"), encoding="utf-8") if x.strip() and not x.startswith("#")]
        return [os.path.join(DIRECT_DIR, n) for n in names if os.path.exists(os.path.join(DIRECT_DIR, n))]
    except OSError:
        return sorted(glob.glob(os.path.join(DIRECT_DIR, "*.tsv")))


def campaign_urls(paths: Optional[List[str]] = None) -> Dict[str, List[str]]:
    """{страница без query: [якоря]}"""
    out: Dict[str, set] = {}
    for p in paths or active_files():
        for r in read_tsv(p):
            for u in [r.get("Ссылка", "")] + r.get("Адреса быстрых ссылок", "").split("||"):
                u = u.strip()
                if not u.startswith("http"):
                    continue
                base, _, frag = u.partition("#")
                page = base.split("?")[0]
                out.setdefault(page, set())
                if frag:
                    out[page].add(frag)
    return {k: sorted(v) for k, v in out.items()}


def check_page(url: str, anchors: List[str]) -> List[str]:
    probs = []
    # Тильда иногда отдаёт роботам (запросы с серверов GitHub) разовый 402/5xx, хотя подписка оплачена
    # и в браузере страница открывается. Поэтому повторяем 3 раза с паузой и ругаемся, только если не прошло.
    r = None
    for i in range(3):
        try:
            r = requests.get(url, headers=UA, timeout=30, allow_redirects=True)
        except requests.RequestException as exc:
            if i == 2:
                return [f"не открывается: {exc.__class__.__name__}"]
            time.sleep(20)
            continue
        if r.status_code == 200 or r.status_code == 404:
            break
        time.sleep(20)
    if r.status_code != 200:
        return [f"ответ {r.status_code}" + (" (Тильда 3 раза подряд не отдала страницу — проверьте её в браузере)" if r.status_code == 402 else "")]
    html = r.text
    low = html.lower()
    for ph in PLACEHOLDERS:
        if ph in low:
            probs.append(f"заглушка шаблона «{ph}»")
    for a in anchors:
        if a.startswith("rec") or a in ("zeropopup", "opencart", "order"):
            continue
        if not re.search(r'(id|name|data-anchor|data-record-anchor)\s*=\s*["\']%s["\']' % re.escape(a), html, re.I) \
                and f'href="#{a}"' not in html:
            probs.append(f"нет якоря #{a}")
    return probs


def cmd_direct_check_links(paths: Optional[List[str]] = None, notify: bool = True) -> int:
    urls = campaign_urls(paths)
    if not urls:
        print("Нет ссылок в data/direct/*.tsv")
        return 0
    lines = []
    for page, anchors in sorted(urls.items()):
        probs = check_page(page, anchors)
        print(("OK  " if not probs else "ПЛОХО ") + page + ("" if not probs else ": " + "; ".join(probs)))
        if probs:
            lines.append(f"• {page.replace('https://', '')}: " + "; ".join(probs))
    st = _state()
    digest = hashlib.md5("\n".join(lines).encode()).hexdigest() if lines else ""
    today = datetime.utcnow().strftime("%Y-%m-%d")
    if lines and notify and (st.get("links_digest") != digest or st.get("links_day") != today):
        if st.get("links_digest") != digest:
            _tg_owner("⚠️ Директ: реклама ведёт на страницы с проблемами — клики будут теряться:\n\n" + "\n".join(lines)
                      + "\n\nНапишите Claude — поправим ссылки кампании или страницу.")
        st["links_day"] = today
    elif not lines and st.get("links_digest"):
        _tg_owner("✅ Директ: все страницы из кампании снова открываются нормально.")
    st["links_digest"] = digest
    _save_state(st)
    return 0


# ---------------------------------------------------------------- сводка из API Директа

REPORTS_URL = "https://api.direct.yandex.com/json/v5/reports"


def _report(token: str, body: dict, tries: int = 12) -> List[Dict[str, str]]:
    headers = {"Authorization": "Bearer " + token, "Accept-Language": "ru", "processingMode": "auto",
               "returnMoneyInMicros": "false", "skipReportHeader": "true", "skipReportSummary": "true"}
    login = os.getenv("YANDEX_DIRECT_LOGIN", "")
    if login:
        headers["Client-Login"] = login
    for _ in range(tries):
        r = requests.post(REPORTS_URL, json=body, headers=headers, timeout=120)
        if r.status_code == 200:
            rows = list(csv.reader(r.text.splitlines(), delimiter="\t"))
            if not rows:
                return []
            return [dict(zip(rows[0], x)) for x in rows[1:]]
        if r.status_code in (201, 202):
            time.sleep(min(int(r.headers.get("retryIn", "10") or 10), 30))
            continue
        raise RuntimeError(f"Директ API {r.status_code}: {r.text[:300]}")
    raise RuntimeError("Директ API: отчёт не успел сформироваться")


def _num(x: str) -> float:
    try:
        return float(str(x).replace(",", ".")) if x not in ("", "--") else 0.0
    except ValueError:
        return 0.0


def _perf(token: str, d_from: date, d_to: date, name: str) -> Dict[str, float]:
    body = {"params": {
        "SelectionCriteria": {"DateFrom": d_from.isoformat(), "DateTo": d_to.isoformat()},
        "FieldNames": ["Impressions", "Clicks", "Cost", "Conversions"],
        "ReportName": f"{name}-{d_from}-{d_to}-{int(time.time())}",
        "ReportType": "ACCOUNT_PERFORMANCE_REPORT", "DateRangeType": "CUSTOM_DATE",
        "Format": "TSV", "IncludeVAT": "YES", "IncludeDiscount": "NO"}}
    rows = _report(token, body)
    tot = {"Impressions": 0.0, "Clicks": 0.0, "Cost": 0.0, "Conversions": 0.0}
    for r in rows:
        for k in tot:
            tot[k] += _num(r.get(k, "0"))
    return tot


def _fmt(t: Dict[str, float]) -> str:
    cl = t["Clicks"]
    cpc = t["Cost"] / cl if cl else 0
    conv = t["Conversions"]
    s = f"расход {t['Cost']:,.0f} ₽ · показы {t['Impressions']:,.0f} · клики {cl:,.0f} · клик {cpc:,.0f} ₽"
    s += f" · конверсии {conv:,.0f}" + (f" (заявка {t['Cost'] / conv:,.0f} ₽)" if conv else "")
    return s.replace(",", " ")


def cmd_direct_report() -> int:
    token = os.getenv("YANDEX_DIRECT_TOKEN", "")
    if not token:
        print("YANDEX_DIRECT_TOKEN не задан — сводку Директа пропускаю.")
        return 0
    y = date.today() - timedelta(days=1)
    week = _perf(token, y - timedelta(days=6), y, "acc7")
    day = _perf(token, y, y, "acc1")
    if not week["Impressions"] and not day["Impressions"]:
        print("Показов за 7 дней нет — сводку не отправляю.")
        return 0
    lines = [f"📊 Яндекс Директ, {y:%d.%m} (с НДС)", "Вчера: " + _fmt(day), "7 дней: " + _fmt(week)]
    try:
        q = _report(token, {"params": {
            "SelectionCriteria": {"DateFrom": y.isoformat(), "DateTo": y.isoformat()},
            "FieldNames": ["Query", "Clicks", "Cost"],
            "ReportName": f"sq-{y}-{int(time.time())}", "ReportType": "SEARCH_QUERY_PERFORMANCE_REPORT",
            "DateRangeType": "CUSTOM_DATE", "Format": "TSV", "IncludeVAT": "YES"}})
        q = sorted(q, key=lambda r: -_num(r.get("Cost", "0")))[:10]
        q = [r for r in q if _num(r.get("Cost", "0")) > 0]
        if q:
            lines.append("\nДороже всего вчера (проверьте, нет ли мусора — пришлите Claude, допишем минус-слова):")
            lines += [f"• {r['Query']} — {_num(r['Cost']):,.0f} ₽, {int(_num(r['Clicks']))} кл.".replace(",", " ") for r in q]
    except Exception as exc:  # noqa — сводка важнее списка запросов
        print(f"Поисковые запросы: {exc}")
    _tg_owner("\n".join(lines))
    print("\n".join(lines))
    return 0


# ---------------------------------------------------------------------------
# Excel-импорт кампаний (Инструменты → «Управление с помощью Excel»), пока нет доступа к API.
TEMPLATE_URL = "https://doc-static.yandex.net/src/support/direct/ru/files/direct_example.xlsx"
TEMPLATE_PATH = os.path.join("data", "direct", "direct_example.xlsx")


def _ensure_template() -> str:
    """Скачать официальный шаблон Директа один раз (из облака Яндекс недоступен — качаем в Actions)."""
    if not os.path.exists(TEMPLATE_PATH):
        r = requests.get(TEMPLATE_URL, timeout=60)
        r.raise_for_status()
        with open(TEMPLATE_PATH, "wb") as f:
            f.write(r.content)
        print(f"Шаблон скачан: {TEMPLATE_PATH} ({len(r.content)} байт)")
    return TEMPLATE_PATH


def cmd_direct_template_dump() -> int:
    """Скачать шаблон и вывести его структуру (листы, первые строки) в лог — чтобы собрать файлы точно по нему."""
    import openpyxl

    path = _ensure_template()
    wb = openpyxl.load_workbook(path)
    for ws in wb.worksheets:
        print(f"\n### Лист «{ws.title}»: {ws.max_row} строк × {ws.max_column} колонок; объединения: {list(ws.merged_cells.ranges)[:20]}")
        for i, row in enumerate(ws.iter_rows(max_row=min(ws.max_row, 40), values_only=True), 1):
            vals = ["" if v is None else str(v) for v in row]
            while vals and vals[-1] == "":
                vals.pop()
            if vals:
                print(f"{i:>3}: " + " | ".join(vals))
    return 0
