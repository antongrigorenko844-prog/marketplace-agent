"""
Заполнение характеристик (атрибутов) карточек Ozon — для контент-рейтинга.

Два шага, оба безопасные:
  ozon-attr-plan   (ТОЛЬКО ЧТЕНИЕ) — для артикулов из data/ozon_attr_articles.txt читает
                   текущие характеристики и справочник категории и предлагает значения
                   -> data/ozon_attr_plan.xlsx. Предлагает только то, что точно известно
                   (OEM-номер и др. артикулы из названия, вид техники, количество, хештеги
                   из мастера, комплектация). Остальное — «нужны данные», без догадок.
  ozon-attr-apply  — отправляет ТОЛЬКО строки плана с «да» в колонке «Применить».
                   Метод /v1/product/attributes/update меняет ТОЛЬКО перечисленные
                   характеристики: название, фото, описание, цена, остатки не трогаются.
                   Уже заполненные на Ozon значения НЕ перезаписываются.
                   Не больше MAX_CARDS карточек за запуск. Итог -> data/ozon_attr_result.txt
                   (статус задачи Ozon + рейтинг до/после).
"""
from __future__ import annotations

import json
import os
import re
import time
from typing import Dict, List, Optional

DATA = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data")
ARTICLES = os.path.join(DATA, "ozon_attr_articles.txt")
PLAN = os.path.join(DATA, "ozon_attr_plan.xlsx")
RESULT = os.path.join(DATA, "ozon_attr_result.txt")
MAX_CARDS = 10

# Характеристики, которые Ozon чаще всего просит заполнить (по контент-рейтингу 27.10)
TARGET = ["OEM-номер", "Альтернативные артикулы товара", "Вид техники", "Комплектация",
          "Количество в упаковке, шт", "Количество, штук", "Единиц в одном товаре", "#Хештеги",
          "Материал", "Цвет товара", "Расположение детали", "Длина, мм", "Ширина, мм", "Высота, мм",
          "Внутренний диаметр, мм", "Внешний диаметр, мм"]

# OEM-коды VAG/DSG вида 0AM325413A, 02E305051C, 0CW325025H, 01X301127C, WHT001922, 0B5...
CODE_RE = re.compile(r"\b(0[0-9A-Z]{2}\s?\d{3}\s?\d{3}[A-Z]{0,3}|WHT\d{6}|N\d{8,9})\b")

HEAD = ["Артикул", "Название", "Атрибут", "id", "Тип", "Сейчас на Ozon", "Предлагаю",
        "dictionary_value_id", "Источник", "Применить (да)"]


def _articles() -> List[str]:
    if not os.path.exists(ARTICLES):
        return []
    out = []
    for line in open(ARTICLES, encoding="utf-8"):
        line = line.split("#", 1)[0].strip()
        out += [a.strip() for a in line.split(",") if a.strip()]
    return out


def _master() -> Dict[str, dict]:
    import openpyxl

    ws = openpyxl.load_workbook(os.path.join(DATA, "master_control.xlsx"), read_only=True, data_only=True)["Товары"]
    rows = list(ws.iter_rows(values_only=True))
    head = [str(h or "") for h in rows[0]]
    return {str(r[0]).strip(): dict(zip(head, r)) for r in rows[1:] if r[0]}


def _rating_missing() -> Dict[str, List[str]]:
    import openpyxl

    path = os.path.join(DATA, "ozon_content_rating.xlsx")
    if not os.path.exists(path):
        return {}
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True)["Карточки"]
    rows = list(ws.iter_rows(values_only=True))
    return {str(r[0]).strip(): [a.strip() for a in str(r[-1] or "").split(",") if a.strip()] for r in rows[1:] if r[0]}


def _codes(text: str) -> List[str]:
    seen = []
    for m in CODE_RE.findall((text or "").upper()):
        c = m.replace(" ", "")
        if c not in seen:
            seen.append(c)
    return seen


def _qty(name: str) -> Optional[int]:
    m = re.search(r"(\d+)\s*шт", name or "", re.I)
    if m:
        return int(m.group(1))
    if re.search(r"комплект|набор|ремкомплект|\+", name or "", re.I):
        return None  # набор — количество деталей неочевидно, не гадаем
    return 1


def _current(info: dict) -> Dict[int, str]:
    out = {}
    for a in info.get("attributes") or []:
        vals = [str(v.get("value")) for v in a.get("values") or [] if v.get("value") not in (None, "")]
        if vals:
            out[int(a["id"])] = "; ".join(vals)
    return out


def _dict_pick(oc, dcid, tid, attr_id, wanted: str) -> Optional[dict]:
    try:
        vals = oc.search_attribute_values(dcid, tid, attr_id, wanted)
    except Exception:  # noqa: BLE001
        return None
    for v in vals:
        if str(v.get("value", "")).strip().lower() == wanted.lower():
            return v
    return vals[0] if len(vals) == 1 else None


def _dict_options(oc, dcid, tid, attr_id, limit: int = 50) -> List[dict]:
    """Значения справочника характеристики (первые limit) — чтобы выбрать точное, а не угадывать."""
    try:
        data = oc._post("/v1/description-category/attribute/values", {
            "attribute_id": attr_id, "description_category_id": dcid, "type_id": tid,
            "language": "DEFAULT", "last_value_id": 0, "limit": limit})
    except Exception as exc:  # noqa: BLE001
        print(f"  справочник {attr_id}: {exc}")
        return []
    return data.get("result") or []


def cmd_plan() -> int:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    import ozon_client as oc

    arts = _articles()
    if not arts:
        print(f"Нет артикулов: впишите их в {ARTICLES} (через запятую)")
        return 1
    master = _master()
    infos = {str(i.get("offer_id")): i for i in oc.get_product_names(arts)}
    rows = []
    for art in arts:
        info = infos.get(art)
        if not info:
            rows.append([art, "", "— карточка не найдена на Ozon —"] + [""] * 7)
            continue
        dcid, tid = info.get("description_category_id"), info.get("type_id")
        cat = {int(a["id"]): a for a in oc.get_category_attributes(dcid, tid)}
        by_name = {str(a.get("name")).strip(): a for a in cat.values()}
        cur = _current(info)
        m = master.get(art, {})
        name = str(m.get("Название товара") or info.get("name") or "")
        codes = _codes(name) or _codes(art)
        want = [n for n in TARGET if n in by_name]
        for an in want:
            a = by_name.get(an)
            if not a or an in ("Изображения", "Озон.Видео: ссылка", "Озон.Видеообложка: ссылка", "Rich-контент JSON"):
                continue  # медиа и rich — отдельная задача
            aid = int(a["id"])
            typ = ("справочник" if a.get("dictionary_id") else str(a.get("type") or "")) + (" (список)" if a.get("is_collection") else "")
            now = cur.get(aid, "")
            prop, dict_id, src = "", "", "нужны данные от вас"
            if now:
                src = "уже заполнено — не трогаем"
            elif an == "OEM-номер" and codes:
                prop, src = codes[0], "артикул из названия"
            elif an == "Альтернативные артикулы товара":
                alt = codes[1:] + [c for c in _codes(m.get("Описание") or "") if c not in codes]
                if alt:
                    prop, src = ", ".join(alt[:5]), "другие коды из названия/описания"
            elif an == "Вид техники":
                opts = _dict_options(oc, dcid, tid, aid)
                v = next((o for o in opts if "легков" in str(o.get("value", "")).lower()), None)
                if v:
                    prop, dict_id, src = v.get("value"), v.get("id"), "справочник Ozon"
                else:
                    src = "варианты: " + "; ".join(f"{o.get('value')} [{o.get('id')}]" for o in opts[:20])
            elif an in ("Количество в упаковке, шт", "Количество, штук", "Единиц в одном товаре"):
                q = _qty(name)
                if q:
                    prop, src = str(q), "из названия" if q > 1 else "одна деталь"
            elif an == "#Хештеги":
                tags = str(m.get("Хэштеги / Теги") or "").strip()
                if tags:
                    prop, src = tags, "колонка «Хэштеги» мастера"
            elif an == "Комплектация":
                q = _qty(name)
                base = re.split(r"\s+(мехатроник|DSG|DQ\d)", name, maxsplit=1)[0].strip()
                if q and base:
                    prop, src = f"{base} — {q} шт.", "из названия — проверьте"
            if not now and not prop and a.get("dictionary_id") and an != "Вид техники":
                opts = _dict_options(oc, dcid, tid, aid)
                if opts:
                    src = "выберите: " + "; ".join(f"{o.get('value')} [{o.get('id')}]" for o in opts[:25])
            rows.append([art, name, an, aid, typ, now, prop, dict_id, src, ""])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "План"
    ws.append(HEAD)
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="305496")
    for i, w in enumerate([22, 40, 30, 8, 14, 25, 40, 12, 28, 12], 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        if row[6].value:
            row[6].fill = PatternFill("solid", fgColor="C6EFCE")
    wb.save(PLAN)
    ready = sum(1 for r in rows if r[6])
    print(f"План: {len(rows)} строк, готовых значений {ready}; сохранено {PLAN}")
    for r in rows:
        print(f"  {r[0]} | {r[2]} | сейчас: {r[5] or '—'} | предлагаю: {r[6] or '—'} | {r[8]}")
    return 0


def _rating(oc, skus: List[int]) -> Dict[int, float]:
    if not skus:
        return {}
    try:
        data = oc._post("/v1/product/rating-by-sku", {"skus": skus})
    except Exception:  # noqa: BLE001
        return {}
    return {int(p["sku"]): p.get("rating") for p in data.get("products") or []}


def cmd_apply() -> int:
    import openpyxl

    import ozon_client as oc

    if not os.path.exists(PLAN):
        print("Нет плана — сначала ozon-attr-plan")
        return 1
    ws = openpyxl.load_workbook(PLAN, data_only=True)["План"]
    rows = list(ws.iter_rows(values_only=True))
    idx = {h: i for i, h in enumerate(rows[0])}
    plan: Dict[str, List[tuple]] = {}
    for r in rows[1:]:
        if str(r[idx["Применить (да)"]] or "").strip().lower() != "да" or not r[idx["Предлагаю"]]:
            continue
        plan.setdefault(str(r[0]).strip(), []).append(r)
    if not plan:
        print("В плане нет строк с «да» — отправлять нечего.")
        return 0
    if len(plan) > MAX_CARDS:
        print(f"Отказ: {len(plan)} карточек за раз, лимит {MAX_CARDS}. Разбейте план.")
        return 1

    infos = {str(i.get("offer_id")): i for i in oc.get_product_names(list(plan))}
    log = []
    items = []
    for art, prs in plan.items():
        cur = _current(infos.get(art) or {})
        attrs = []
        for r in prs:
            aid = int(r[idx["id"]])
            if cur.get(aid):
                log.append(f"{art}: {r[idx['Атрибут']]} уже заполнено на Ozon ({cur[aid]}) — пропускаю")
                continue
            val, did = str(r[idx["Предлагаю"]]), r[idx["dictionary_value_id"]]
            coll = "(список)" in str(r[idx["Тип"]] or "")
            if did:
                values = [{"dictionary_value_id": int(did), "value": val}]
            elif coll:
                values = [{"value": v.strip()} for v in val.split(",") if v.strip()]
            else:
                values = [{"value": val}]
            attrs.append({"id": aid, "complex_id": 0, "values": values})
            log.append(f"{art}: {r[idx['Атрибут']]} = {val}")
        if attrs:
            items.append({"offer_id": art, "attributes": attrs})

    skus = {}
    for inf in oc.get_product_info_list(list(plan)):
        s = inf.get("sku") or next((x.get("sku") for x in inf.get("sources") or [] if x.get("sku")), None)
        if s:
            skus[str(inf.get("offer_id"))] = int(s)
    before = _rating(oc, list(skus.values()))

    if not items:
        log.append("Отправлять нечего.")
    else:
        resp = oc.update_attributes(items)
        task = resp.get("task_id") or (resp.get("result") or {}).get("task_id")
        log.append(f"Отправлено карточек: {len(items)}; task_id={task}; ответ: {json.dumps(resp, ensure_ascii=False)[:300]}")
        if task:
            for _ in range(12):
                time.sleep(10)
                st = oc.get_import_status(int(task))
                its = (st.get("result") or {}).get("items") or st.get("items") or []
                if its and all(str(i.get("status")) not in ("pending", "processing", "") for i in its):
                    break
            for i in its:
                log.append(f"  Ozon: {i.get('offer_id')} статус {i.get('status')} ошибки {json.dumps(i.get('errors') or [], ensure_ascii=False)[:400]}")
        time.sleep(20)
        after = _rating(oc, list(skus.values()))
        for art, s in skus.items():
            log.append(f"  рейтинг {art}: {before.get(s)} -> {after.get(s)} (Ozon может пересчитать позже)")
    open(RESULT, "w", encoding="utf-8").write("\n".join(log) + "\n")
    print("\n".join(log))
    return 0
