"""
Единая точка добавления СОВСЕМ НОВОГО товара сразу на Ozon + WB (а после
короткого ожидания — автоматически и на Avito + сайт, т.к. их фиды строятся
из общего data/ozon_catalog.xlsx).

Раньше нужно было отдельно заполнять data/ozon_new_products.xlsx и
data/wb_new_products.xlsx — почти теми же данными дважды, плюс отдельно
подставлять фото. Здесь ОДНА таблица data/new_products.xlsx: артикул
используется одновременно как offer_id на Ozon и vendorCode на WB (в
проекте они и так всегда совпадают — см. --compare-ozon-wb). "Образец"
обычно один на обе площадки; если у нового товара образец продаётся
ТОЛЬКО на Ozon (на WB такого раздела/линейки ещё нет) — заполните
отдельную колонку "Образец ТОЛЬКО для WB" другим, похожим по типу
товаром, который уже есть на WB (раздел/характеристики/габариты
возьмутся у него, остальное всё равно из своих данных). Фото берутся
автоматически из photos/<артикул>/
(как для всех остальных товаров) — заливать их в таблицу вручную не нужно.

ЧЕСТНО О ГРАНИЦАХ (это ограничения самих Ozon/WB API, обойти нельзя):
  - Ozon создаёт карточку сразу с фото за один вызов.
  - WB создаёт карточку БЕЗ фото — у метода создания нет для этого
    возможности (нужен nmID, а WB не отдаёт его сразу при создании). Фото
    на WB добавляются вторым шагом, через 2-5 минут, когда nmID уже
    известен — см. main.py --finish-new-product.
  - Avito и сайт (Тильда) строят фид ИЗ ozon_catalog.xlsx — то есть видят
    только то, что уже реально подтверждено на Ozon. Нужен ещё один шаг
    (fetch-ozon + build-ozon-catalog) после того, как карточка на Ozon
    реально создалась — тоже входит в --finish-new-product.

Порядок команд (см. main.py):
  1. build-new-product-template — data/new_products.xlsx.
  2. Заполнить строку (артикул, образец, название, описание, цена...);
     положить фото в photos/<артикул>/.
  3. push-new-product-all-dryrun — посмотреть, что будет отправлено.
  4. push-new-product-all — реально создать на Ozon и отправить на создание
     на WB (WB — без фото, это нормально на этом шаге).
  5. Подождать 2-5 минут.
  6. finish-new-product — подтягивает новые товары в общий каталог,
     доливает фото на WB, обновляет фиды Avito и сайта.
"""
import logging
import os
import re
from typing import Dict, List

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

logger = logging.getLogger("marketplace-agent.new_product")

_EMOJI_RE = re.compile(
    "["
    "\U0001F300-\U0001FAFF"  # пиктограммы (🔧📦🔩 и т.п.)
    "\U00002600-\U000027BF"  # значки/дингбаты (✅ и т.п.)
    "\U00002B00-\U00002BFF"  # доп. символы/стрелки (⭐ и т.п.)
    "\U0000FE0F"              # variation selector, часто идёт следом за эмодзи
    "]+"
)


def _strip_emoji_for_wb(text):
    """
    WB (в отличие от Ozon) реально отклоняет создание карточки, если в
    описании есть эмодзи — ошибка "Поле Описание не должно содержать
    запрещенные символы: ...". Мы используем эмодзи в описаниях для Ozon
    осознанно (разбивка на смысловые блоки для SEO), поэтому не убираем
    их из общего edits[...]['description'] — а только из WB-версии здесь,
    прямо перед отправкой на WB.
    """
    if not text:
        return text
    cleaned = _EMOJI_RE.sub("", text)
    cleaned = re.sub(r" {2,}", " ", cleaned)
    cleaned = re.sub(r"\n[ \t]+", "\n", cleaned)
    return cleaned.strip()


FONT_NAME = "Arial"
HEADER_FILL = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF")

COLUMNS = [
    ("offer_id", "НОВЫЙ артикул — один и тот же для Ozon и WB, придумайте сами"),
    ("sample_offer_id", "Образец: артикул похожего товара, уже продающегося и на Ozon, и на WB"),
    ("sample_offer_id_wb", "Образец ТОЛЬКО для WB, если Ozon-образец на WB не продаётся (необязательно — пусто = взять тот же, что для Ozon)"),
    ("name", "Название (для WB автоматически обрежется до 60 символов)"),
    ("description", "Описание"),
    ("price", "Цена, ₽"),
    ("old_price", "Цена до скидки, ₽ (необязательно, только Ozon)"),
    ("tnved", "Код ТН ВЭД — обязательно для Ozon (иначе карточка останется черновиком)"),
    ("weight_g", "Вес, г (пусто = как у образца)"),
    ("length_mm", "Длина, мм (пусто = как у образца)"),
    ("width_mm", "Ширина, мм (пусто = как у образца)"),
    ("height_mm", "Высота, мм (пусто = как у образца)"),
    ("quantity_to_sell", "Кол-во к продаже (остаток)"),
    ("notes", "Заметки"),
    ("hashtags", "Хэштеги / ключевые слова через запятую (необязательно, только Ozon)"),
]
WIDTHS = [20, 30, 30, 45, 45, 12, 18, 20, 12, 12, 12, 12, 16, 25, 30]


def build_template(xlsx_path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Новый товар — все площадки"

    for col_idx, (_, header) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL

    for col_idx, width in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}1"
    os.makedirs(os.path.dirname(xlsx_path), exist_ok=True)
    wb.save(xlsx_path)


def load_edits(xlsx_path: str) -> Dict[str, dict]:
    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb.active
    edits: Dict[str, dict] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] in (None, ""):
            continue
        offer_id = str(row[0]).strip()
        raw = {COLUMNS[i][0]: (row[i] if i < len(row) else None) for i in range(len(COLUMNS))}
        edits[offer_id] = {
            "sample_offer_id": str(raw.get("sample_offer_id") or "").strip(),
            "sample_offer_id_wb": str(raw.get("sample_offer_id_wb") or "").strip(),
            "name": (raw.get("name") or "").strip() if isinstance(raw.get("name"), str) else raw.get("name"),
            "description": raw.get("description") or "",
            "price": raw.get("price"),
            "old_price": raw.get("old_price"),
            "tnved": str(raw.get("tnved") or "").strip(),
            "weight_g": raw.get("weight_g"),
            "length_mm": raw.get("length_mm"),
            "width_mm": raw.get("width_mm"),
            "height_mm": raw.get("height_mm"),
            "quantity_to_sell": raw.get("quantity_to_sell"),
            "notes": raw.get("notes"),
            "hashtags": (raw.get("hashtags") or "").strip() if isinstance(raw.get("hashtags"), str) else (raw.get("hashtags") or ""),
            "images": [],
        }
    return edits


def attach_photos(edits: Dict[str, dict], photos_dir: str) -> Dict[str, List[str]]:
    """
    Заливает фото из photos/<offer_id>/ (та же папка и та же логика, что и
    для всех остальных товаров) на GitHub Release через photo_host.py и
    дописывает ссылки прямо в edits[offer_id]["images"] — отдельный шаг
    attach-*-new-photos тут не нужен, всё в одном push-new-product-all*.
    """
    import catalog_editor
    import photo_host

    matched: Dict[str, List[str]] = {}
    for offer_id, edit in edits.items():
        own_files = catalog_editor.own_media_files(photos_dir, offer_id, catalog_editor.IMAGE_EXTS)
        if not own_files:
            continue
        urls = []
        for f in own_files:
            try:
                url = photo_host.upload_file(
                    catalog_editor.media_path(photos_dir, offer_id, f),
                    filename=catalog_editor.asset_filename_for(offer_id, f),
                )
            except Exception as exc:
                logger.warning("%s: не удалось загрузить фото %s: %s", offer_id, f, exc)
                continue
            urls.append(url)
        if urls:
            edit["images"] = urls
            matched[offer_id] = urls
    return matched


def to_ozon_edits(edits: Dict[str, dict]) -> Dict[str, dict]:
    """Формат, ожидаемый ozon_new_products.build_new_import_items."""
    out: Dict[str, dict] = {}
    for offer_id, e in edits.items():
        out[offer_id] = {
            "sample_offer_id": e.get("sample_offer_id"),
            "name": e.get("name"),
            "description": e.get("description"),
            "price": e.get("price"),
            "old_price": e.get("old_price"),
            "barcode": "",
            "tnved": e.get("tnved"),
            "weight_g": e.get("weight_g"),
            "length_mm": e.get("length_mm"),
            "width_mm": e.get("width_mm"),
            "height_mm": e.get("height_mm"),
            "images": e.get("images") or [],
            "video_url": "",
            "quantity_to_sell": e.get("quantity_to_sell"),
            "notes": e.get("notes"),
            "hashtags": e.get("hashtags"),
        }
    return out


def _mm_to_cm(value):
    if value in (None, ""):
        return None
    try:
        return float(value) / 10
    except (TypeError, ValueError):
        return None


def _g_to_kg(value):
    if value in (None, ""):
        return None
    try:
        return float(value) / 1000
    except (TypeError, ValueError):
        return None


def to_wb_edits(edits: Dict[str, dict]) -> Dict[str, dict]:
    """Формат, ожидаемый wb_new_products.build_new_card_groups (см. там же про единицы: кг/см)."""
    out: Dict[str, dict] = {}
    for offer_id, e in edits.items():
        name = e.get("name") or ""
        out[offer_id] = {
            "sample_vendor_code": e.get("sample_offer_id_wb") or e.get("sample_offer_id"),
            "title": _strip_emoji_for_wb(name)[:60],
            "description": _strip_emoji_for_wb(e.get("description")),
            "price": e.get("price"),
            "weight_kg": _g_to_kg(e.get("weight_g")),
            "length_cm": _mm_to_cm(e.get("length_mm")),
            "width_cm": _mm_to_cm(e.get("width_mm")),
            "height_cm": _mm_to_cm(e.get("height_mm")),
            "notes": e.get("notes"),
        }
    return out
