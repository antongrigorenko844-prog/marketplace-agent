"""
Вес и габариты упаковки из data/master_control.xlsx (колонки «Вес, кг»,
«Длина, см», «Ширина, см», «Высота, см») — ОДИН источник для Ozon и WB.

Используется:
  - catalog_editor.py  — при отправке карточки Ozon (вес 4497 + depth/width/height);
  - wb_catalog_editor.py — при отправке карточки WB (dimensions);
  - validate_master.py — предупреждения о неполно заполненных габаритах.

Правило: габариты берутся, только если заполнены ВСЕ четыре колонки. Пусто —
значит «не трогать», и на площадке остаётся то, что там уже стоит.
Старые записи weight_g / dims_mm в data/ozon_overrides.json по-прежнему
работают, но мастер, если в нём заполнено, главнее.
"""
import os
from typing import Dict, List, Optional

from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.abspath(__file__))
MASTER_PATH = os.path.join(ROOT, "data", "master_control.xlsx")

DIM_COLUMNS = ["Вес, кг", "Длина, см", "Ширина, см", "Высота, см"]


def _num(v) -> Optional[float]:
    if v is None or str(v).strip() == "":
        return None
    try:
        return float(str(v).replace(" ", "").replace("\xa0", "").replace(",", "."))
    except ValueError:
        return None


def read_rows(path: str = MASTER_PATH) -> Dict[str, List[Optional[float]]]:
    """{артикул: [вес_кг, длина_см, ширина_см, высота_см]} — как в мастере, с пустыми (None)."""
    if not os.path.exists(path):
        return {}
    try:
        ws = load_workbook(path, read_only=True, data_only=True)["Товары"]
    except Exception:
        return {}
    rows = ws.iter_rows(values_only=True)
    header = [str(h or "").strip() for h in next(rows, [])]
    if not all(c in header for c in DIM_COLUMNS):
        return {}
    idx = [header.index(c) for c in DIM_COLUMNS]
    out: Dict[str, List[Optional[float]]] = {}
    for r in rows:
        if not r or r[0] is None or str(r[0]).strip() == "":
            continue
        vals = [_num(r[i]) if i < len(r) else None for i in idx]
        if any(v is not None for v in vals):
            out[str(r[0]).strip()] = vals
    return out


_CACHE: Dict[tuple, Dict[str, dict]] = {}


def load(path: str = MASTER_PATH) -> Dict[str, dict]:
    """{артикул: {"weight_g": int, "dims_mm": [L, W, H]}} — только полностью заполненные строки.
    Результат кэшируется по времени изменения файла (вызывается на каждую карточку)."""
    try:
        stamp = (path, os.path.getmtime(path))
    except OSError:
        return {}
    if stamp in _CACHE:
        return _CACHE[stamp]
    out: Dict[str, dict] = {}
    for art, (w, l, wd, h) in read_rows(path).items():
        if None in (w, l, wd, h) or min(w, l, wd, h) <= 0:
            continue
        out[art] = {
            "weight_g": int(round(w * 1000)),
            "dims_mm": [int(round(l * 10)), int(round(wd * 10)), int(round(h * 10))],
        }
    _CACHE.clear()
    _CACHE[stamp] = out
    return out


def for_article(article: str, cache: Optional[Dict[str, dict]] = None) -> Optional[dict]:
    """Габариты по артикулу (точное совпадение, затем без учёта регистра)."""
    data = cache if cache is not None else load()
    if article in data:
        return data[article]
    low = str(article).lower()
    for k, v in data.items():
        if k.lower() == low:
            return v
    return None
